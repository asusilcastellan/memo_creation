"""The analyst's review list: "<memo> - Review.xlsx", saved next to the memo.

    build_review_workbook(path, facts, flags, plan, extracted, usage_summary_for_internal_log_only=None)

Sheets: Review items (one row per facts.Flag, same IDs as the red remarks in the .docx) | Online search |
Background checks | Sources used | Key figures | Email chain (only when emails were read).
Build the .docx FIRST: it assigns the IDs. Nothing technical goes in here: no token counts, model names or
file-system paths outside the deal folder (source files are relative to the deal folder).
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "002060"
FILLS = {"review": "FFC7CE", "missing": "FFEB9C", "check": "E7E6E6"}     # light red / light amber / light grey
LEVEL_ORDER = {"review": 0, "missing": 1, "check": 2}
STATUSES = ["Open", "Resolved", "Not applicable"]
CUR, CUR2, NUM, PCT, PCT2, INT = '"$"#,##0', '"$"#,##0.00', "#,##0", "0.0%", "0.00%", "0"
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


# --------------------------------------------------------------------------- helpers
def _header(ws, row: int, headers: list[str], widths: list[int]) -> None:
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w


def _finish(ws, header_row: int, n_cols: int, last_row: int) -> None:
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{max(last_row, header_row)}"


def _status(ws, col: str, first: int, last: int) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True,
                        showDropDown=False, errorTitle="Status", error="Choose Open, Resolved or Not applicable.")
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{max(first, last)}")
    for r in range(first, last + 1):
        ws[f"{col}{r}"] = "Open"


def _id_num(fl) -> int:
    m = re.match(r"R(\d+)", fl.id or "")
    return int(m.group(1)) if m else 10 ** 6


def _naive(dt) -> Optional[datetime]:
    if not dt:
        return None
    try:
        return dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt
    except (ValueError, OSError):
        return dt.replace(tzinfo=None)


def _surname(name: str) -> str:
    words = [w for w in re.findall(r"[A-Za-z]+", name or "") if w.lower() not in ("jr", "sr", "mr", "ms", "mrs", "dr")]
    return words[-1].lower() if words else ""


# --------------------------------------------------------------------------- sheets
def _review_items(ws, f, flags, plan) -> None:
    ws.title = "Review items"
    headers = ["ID", "Priority", "Memo section", "Item", "Details", "Source document", "Source file(s)",
               "Status", "Analyst comment"]
    ws["A1"] = f"{f.address or 'Deal'} - review items - generated {datetime.now():%Y-%m-%d %H:%M}"
    ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    _header(ws, 2, headers, [7, 10, 26, 40, 70, 24, 40, 15, 40])
    rows = sorted(flags, key=lambda fl: (LEVEL_ORDER.get(fl.level, 9), _id_num(fl)))
    r = 2
    for fl in rows:
        r += 1
        row = plan.rows.get(fl.source_key) if plan is not None and fl.source_key else None
        values = [fl.id, fl.level_label, fl.memo_section, fl.title, fl.details,
                  row.display_name if row else "", "\n".join(fl.source_files)]
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP if i in (3, 4, 5, 7) else TOP
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=FILLS.get(fl.level, "FFFFFF"))
    _status(ws, "H", 3, r)
    _finish(ws, 2, len(headers), r)


def _online_search(ws, f) -> None:
    headers = ["Person", "Category", "Severity", "Finding", "Source", "Status", "Analyst comment"]
    _header(ws, 1, headers, [24, 14, 10, 70, 45, 15, 40])
    r = 1
    for fd in f.online_findings:
        r += 1
        for i, v in enumerate([fd.get("person"), fd.get("category"), fd.get("severity"), fd.get("finding")], start=1):
            ws.cell(row=r, column=i, value=v).alignment = WRAP
        url, title = fd.get("source_url"), fd.get("source_title")
        c = ws.cell(row=r, column=5, value=title or url or "")
        if url:
            c.hyperlink = url
            c.style = "Hyperlink"
        c.alignment = WRAP
    if r == 1:
        r = 2
        ws.cell(row=2, column=4, value=(f.online_search_summary or "No adverse findings.")
                if f.online_search_clean is not None else "The online search was not run.")
    _status(ws, "F", 2, r)
    _finish(ws, 1, len(headers), r)


def _role(name: str, f, x) -> str:
    s = _surname(name)
    pfs = {_surname(p.full_name) for p in (x.pfs.sponsors if x.pfs else [])}
    kp = {_surname(n) for n in (x.term_sheet.key_principals if x.term_sheet else [])}
    resume = {_surname(n) for n in (x.resume.people if x.resume else [])}
    if s in pfs or s in resume:
        return "sponsor"
    if s in kp:
        return "guarantor" if (f.personal_guaranty or f.carveout_guaranty or f.completion_guaranty) else "key principal"
    return "not in deal" if (pfs or kp or resume) else "sponsor"


def _background(ws, f, x, plan) -> None:
    headers = ["Person", "Role in deal", "Credit score", "Bankruptcies", "Judgments/liens", "Criminal", "Civil cases",
               "Past due", "Summary", "Source file"]
    _header(ws, 1, headers, [24, 14, 12, 13, 15, 10, 12, 10, 50, 40])
    yn = lambda v: "" if v is None else ("Yes" if v else "No")  # noqa: E731
    row = plan.rows.get("background") if plan is not None else None
    files = "\n".join(e.rel for e in row.files) if row else ""
    r = 1
    for p in (x.background.people if x.background else []):
        r += 1
        role = _role(p.full_name, f, x)
        vals = [p.full_name, role, p.credit_score, yn(p.bankruptcies), yn(p.judgments_or_liens), yn(p.criminal_records),
                yn(p.pending_civil_cases), yn(p.past_due_accounts),
                (p.adverse_summary or "") + (" - not a sponsor/guarantor: not named in the memo" if role == "not in deal" else ""),
                files]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP if i in (9, 10) else TOP
            if i == 3 and isinstance(v, int):
                c.number_format = INT
        if role == "not in deal":
            for i in range(1, len(headers) + 1):
                ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=FILLS["check"])
    _finish(ws, 1, len(headers), r)


def _sources(ws, plan) -> None:
    headers = ["Document", "Used in memo section(s)", "File(s)", "Folder", "Date modified", "Chosen by"]
    _header(ws, 1, headers, [30, 45, 60, 30, 17, 16])
    r = 1
    for row in (plan.rows.values() if plan is not None else []):
        r += 1
        chosen = ("Skipped by analyst" if row.skipped else "Saved choice" if row.saved else
                  "Analyst" if not row.auto else "Automatic" if row.files else "Not found")
        newest = max((e.modified for e in row.files), default=None)
        folder = row.folder if row.files and row.folder not in (None, ".") else ("(deal folder root)" if row.files else "")
        vals = [row.display_name, ", ".join(row.memo_sections), "\n".join(e.rel for e in row.files), folder,
                _naive(newest), chosen]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP if i in (2, 3) else TOP
            if i == 5 and v:
                c.number_format = "yyyy-mm-dd hh:mm"
    _finish(ws, 1, len(headers), r)


def key_figures(f) -> list[tuple[str, object, str, str]]:
    """(figure, value, Excel number format, source document) - numbers stay numbers."""
    rate = lambda v: v / 100 if isinstance(v, (int, float)) else None  # noqa: E731
    rows = [("Castellan loan", f.loan, CUR, "Signed term sheet")]
    if f.is_participation:
        rows += [("Castellan share", rate(f.castellan_share_pct), PCT2, "Signed term sheet / one pager"),
                 ("Total loan", f.total_loan, CUR, "Signed term sheet / one pager"),
                 ("Lead lender", f.lead_lender, "@", "Signed term sheet / one pager")]
    rows += [("Junior loan", f.junior_loan, CUR, "Signed term sheet"),
             ("Total debt", f.total_debt, CUR, "Signed term sheet / underwriting model"),
             ("Purchase price", f.purchase_price, CUR, "Purchase agreement (PSA)"),
             ("Capex budget", f.capex, CUR, "Underwriting model"),
             ("Total cost", f.total_cost, CUR, "Underwriting model"),
             ("Value used", f.value, CUR, "Appraisal"),
             ("Value type used", f.value_label, "@", "Appraisal")]
    rows += [(f"Appraised value - {k}", v, CUR, "Appraisal") for k, v in (f.all_values or {}).items()]
    rows += [("LTV (as printed in the memo)", f.memo_ltv, PCT, "Computed"),
             ("LTV - Castellan loan", f.ltv, PCT, "Computed"),
             ("LTV - total debt", f.ltv_total, PCT, "Computed"),
             ("LTC - Castellan loan", f.ltc, PCT, "Computed"),
             ("LTC - total debt", f.ltc_total, PCT, "Computed"),
             ("Loan per SF", f.loan_psf, CUR2, "Computed / one pager"),
             ("Comps average $/SF", f.comps_psf, CUR2, "Appraisal / one pager"),
             ("Sponsor equity", f.equity, CUR, "Underwriting model"),
             ("Equity % of total cost", f.equity_pct, PCT, "Computed"),
             ("Sponsor liquidity", f.total_liquidity, CUR, "Personal financial statements")]
    for i, s in enumerate(f.credit_scores or [], start=1):
        try:
            rows.append((f"Credit score {i}", int(str(s).lstrip("+")), INT, "Background & credit reports"))
        except ValueError:
            continue
    rows += [("Interest rate index", f.index_name or ("Fixed rate" if f.is_floating is False else None), "@", "Signed term sheet"),
             ("Index rate", rate(f.index_rate), PCT2, f.index_rate_source or ""),
             ("Spread", rate(f.spread), PCT2, "Signed term sheet"),
             ("Floor rate", rate(f.floor), PCT2, "Signed term sheet"),
             ("Fixed rate", rate(f.fixed_rate), PCT2, "Signed term sheet"),
             ("Effective rate at closing", rate(f.effective_rate), PCT2, "Computed")]
    return rows


def _key_figures(ws, f) -> None:
    headers = ["Figure", "Value", "Source document"]
    _header(ws, 1, headers, [32, 20, 36])
    r = 1
    for name, value, fmt, src in key_figures(f):
        if value is None:
            continue
        r += 1
        ws.cell(row=r, column=1, value=name)
        c = ws.cell(row=r, column=2, value=value)
        c.number_format = fmt if fmt != "@" else "General"
        ws.cell(row=r, column=3, value=src)
    _finish(ws, 1, len(headers), r)


def _emails(ws, f, x, flags) -> None:
    headers = ["Date", "From", "Subject", "Source", "Notes"]
    _header(ws, 1, headers, [17, 34, 50, 18, 60])
    r = 1
    for m in x.emails.messages:
        r += 1
        src = m.source.replace("Graph", "Outlook")
        src = re.sub(r"\s*\(.*?\)", "", src)
        note = "" if m.included else "Left out of the summary (length cap)"
        for i, v in enumerate([_naive(m.date), m.sender, m.subject, src, note], start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP if i in (2, 3, 5) else TOP
            if i == 1 and v:
                c.number_format = "yyyy-mm-dd hh:mm"
    for fl in [fl for fl in flags if fl.source_key == "emails"]:
        r += 1
        ws.cell(row=r, column=3, value=f"Conflict with the term sheet ({fl.id})")
        ws.cell(row=r, column=5, value=fl.details).alignment = WRAP
    _finish(ws, 1, len(headers), r)


# --------------------------------------------------------------------------- main
def build_review_workbook(path, facts, flags, plan, extracted, usage_summary_for_internal_log_only=None) -> Path:
    """Write the review workbook. `usage_summary_for_internal_log_only` is accepted for the caller's debug log
    and deliberately NOT written to the workbook."""
    wb = Workbook()
    _review_items(wb.active, facts, flags, plan)
    _online_search(wb.create_sheet("Online search"), facts)
    _background(wb.create_sheet("Background checks"), facts, extracted, plan)
    _sources(wb.create_sheet("Sources used"), plan)
    _key_figures(wb.create_sheet("Key figures"), facts)
    if getattr(extracted, "emails", None) and extracted.emails.messages:
        _emails(wb.create_sheet("Email chain"), facts, extracted, flags)
    wb.active = 0
    path = Path(path)
    wb.save(path)
    return path
