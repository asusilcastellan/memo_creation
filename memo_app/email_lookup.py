"""Find a deal's Outlook conversation ID from the team's board-items workbook.

The analyst board is exported to an .xlsx (in Dropbox) with one row per deal and
its Outlook 'Conversation ID'. Instead of asking the analyst to paste a conversation
ID, we take the deal folder's final name (e.g. ".../42-18 147th Street" -> "42-18
147th Street"), fuzzy-match it against the configured columns ('Name' and
'Property Address' by default), and let the analyst pick the right row in a dialog.
The chosen row's conversation ID is then read through Graph (graph_mail.py).

Matching uses a partial (substring-aware) ratio so a folder named just for the
street still matches the full "street, city, STATE zip" address in the workbook.

All configuration lives under `email_lookup:` in settings.yaml; nothing here raises
to the caller except load_rows() when the workbook cannot be read (the GUI shows that
message and unticks the checkbox).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Callable

Log = Callable[[str], None]

# Defaults match demo_data.xlsx; override any of these under email_lookup: in settings.yaml.
DEFAULT_SHEET = "board_items"
DEFAULT_MATCH_COLUMNS = ["Name", "Property Address"]
DEFAULT_ID_COLUMN = "Conversation ID"
DEFAULT_DISPLAY_COLUMNS = ["Name", "Broker", "Loan Request", "Loan Type", "Asset Type"]
DEFAULT_THRESHOLD = 80


@dataclass
class LookupRow:
    conversation_id: str
    values: dict[str, str] = field(default_factory=dict)   # display column -> text (for the chooser table)
    score: int = 0                                          # 0-100 match against the current query


@dataclass
class LookupConfig:
    workbook: str = ""
    sheet: str = DEFAULT_SHEET
    match_columns: list[str] = field(default_factory=lambda: list(DEFAULT_MATCH_COLUMNS))
    id_column: str = DEFAULT_ID_COLUMN
    display_columns: list[str] = field(default_factory=lambda: list(DEFAULT_DISPLAY_COLUMNS))
    threshold: int = DEFAULT_THRESHOLD

    @classmethod
    def from_settings(cls, settings) -> "LookupConfig":
        g = settings.get("email_lookup", {}) if settings is not None else {}
        g = g if isinstance(g, dict) else {}
        return cls(
            workbook=str(g.get("workbook") or ""),
            sheet=str(g.get("sheet") or DEFAULT_SHEET),
            match_columns=list(g.get("match_columns") or DEFAULT_MATCH_COLUMNS),
            id_column=str(g.get("id_column") or DEFAULT_ID_COLUMN),
            display_columns=list(g.get("display_columns") or DEFAULT_DISPLAY_COLUMNS),
            threshold=int(g.get("threshold") or DEFAULT_THRESHOLD),
        )


def deal_name_from_path(path: str) -> str:
    """The final folder name of a deal path (local or Dropbox, either separator)."""
    p = (path or "").strip().strip('"').replace("\\", "/").rstrip("/")
    return p.rsplit("/", 1)[-1] if p else ""


def _norm(s) -> str:
    return "" if s is None else str(s).strip().lower()


def partial_ratio(query: str, candidate: str) -> float:
    """difflib ratio of `query` against the best-matching same-length window of `candidate`.

    Substring-aware: "42-18 147th street" vs "42-18 147th street, flushing, ny 11355" -> 1.0.
    """
    a, b = _norm(query), _norm(candidate)
    if not a or not b:
        return 0.0
    short, long = (a, b) if len(a) <= len(b) else (b, a)
    n = len(short)
    if n == len(long):
        return SequenceMatcher(None, short, long).ratio()
    best = 0.0
    for i in range(len(long) - n + 1):
        best = max(best, SequenceMatcher(None, short, long[i:i + n]).ratio())
        if best >= 1.0:
            break
    return best


def score_row(query: str, row: LookupRow, match_columns: list[str]) -> int:
    best = 0.0
    for col in match_columns:
        best = max(best, partial_ratio(query, row.values.get(col, "")))
    return round(best * 100)


def match(rows: list[LookupRow], query: str, cfg: LookupConfig, limit: int = 25) -> list[LookupRow]:
    """Rows scored against `query`, highest first. Returns matches at/above the threshold;
    if none reach it, returns the best `min(limit, 15)` anyway so a near-miss stays pickable."""
    q = (query or "").strip()
    if not q:
        return []
    scored: list[LookupRow] = []
    for r in rows:
        r.score = score_row(q, r, cfg.match_columns)
        scored.append(r)
    scored.sort(key=lambda r: r.score, reverse=True)
    above = [r for r in scored if r.score >= cfg.threshold]
    if above:
        return above[:limit]
    return scored[:min(limit, 15)]


def _resolve_workbook(cfg: LookupConfig, settings, log: Log) -> Path:
    """Local/synced path -> use it directly; otherwise download the Dropbox path to the cache."""
    if not cfg.workbook:
        raise RuntimeError(
            "No lookup workbook configured. Set email_lookup.workbook in settings.yaml to the "
            "board export (a local path or a Dropbox path like /Bridge Lending/.../board.xlsx)."
        )
    local = Path(cfg.workbook).expanduser()
    if local.is_file():
        return local
    if not getattr(settings, "has_dropbox", False):
        raise RuntimeError(
            f"Lookup workbook '{cfg.workbook}' is not a file on this computer and no Dropbox "
            "credentials are configured to download it."
        )
    from .storage import DropboxStorage

    return DropboxStorage(settings, log).download_path(cfg.workbook)


def load_rows(settings, log: Log = print, cfg: LookupConfig | None = None) -> list[LookupRow]:
    """Every deal row that has a conversation ID, with its display + match column values.
    Raises RuntimeError (plain-English) when the workbook cannot be read."""
    import openpyxl

    cfg = cfg or LookupConfig.from_settings(settings)
    path = _resolve_workbook(cfg, settings, log)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise RuntimeError(f"Could not open the lookup workbook: {exc}") from exc
    try:
        ws = wb[cfg.sheet] if cfg.sheet in wb.sheetnames else wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = [("" if c is None else str(c).strip()) for c in next(rows_iter)]
        except StopIteration:
            return []
        idx = {name: i for i, name in enumerate(header)}
        if cfg.id_column not in idx:
            raise RuntimeError(
                f"The lookup workbook has no '{cfg.id_column}' column (sheet '{ws.title}'). "
                "Set email_lookup.id_column in settings.yaml to the right column name."
            )
        wanted = list(dict.fromkeys(cfg.display_columns + cfg.match_columns))

        def cell(r, name):
            i = idx.get(name)
            if i is None or i >= len(r) or r[i] is None:
                return ""
            v = r[i]
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            return str(v).strip()

        out: list[LookupRow] = []
        for r in rows_iter:
            cid = cell(r, cfg.id_column)
            if not cid:
                continue
            out.append(LookupRow(conversation_id=cid, values={name: cell(r, name) for name in wanted}))
        return out
    finally:
        wb.close()
