"""The "Sources" table: which file(s) feed each part of the memo, editable by the analyst before generating.

    build_plan(deal_path, selections, entries, rules)  -> SourcesPlan (from the locator's automatic picks)
    plan.set_files / skip / reset                       analyst overrides (the "Choose source" dialog)
    plan.save() / plan.apply_saved()                    choices persist per deal in cache/plans/<deal>.json
    plan.selections()                                   -> {key: Selection} for the extractor
    pick_sheets(xlsx, deal_name)                        one-pager sheets for THIS property only
"""
from __future__ import annotations

import fnmatch
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path, PurePosixPath
from typing import Any, Optional

from . import cache
from .locator import Selection
from .storage import FileEntry

ROOT_LABEL = "(deal folder root)"

# key -> (display name, used for). sources.yaml may override with display_name / used_for.
DEFAULT_NAMES = {
    "term_sheet": ("Signed term sheet", "loan amount, rate, guaranty"),
    "psa": ("Purchase agreement (PSA)", "purchase price"),
    "model": ("Underwriting model", "budget, sources & uses, property facts"),
    "one_pager": ("One pager", "loan financials tables"),
    "resume": ("Sponsor resume / bio", "business background, experience"),
    "pfs": ("Personal financial statements", "net worth, liquidity"),
    "background": ("Background & credit reports", "background and credit check"),
    "appraisal": ("Appraisal", "value, comps, submarket"),
    "environmental": ("Environmental report(s)", "environmental risk"),
    "insurance": ("Insurance", "insurance risk"),
    "emails": ("Emails", "loan purpose, business plan, exit"),
    "photos": ("Property photos", "pictures"),
}


@dataclass
class SourceRow:
    key: str
    display_name: str
    used_for: str
    multi: bool
    folder: Optional[str]
    files: list[FileEntry]
    memo_sections: list[str] = field(default_factory=list)   # memo sections this document feeds (sources.yaml)
    auto: bool = True                       # True = picked by the rules in sources.yaml
    sheets: list[str] = field(default_factory=list)    # one pager only: the sheets whose tables go in the memo
    excluded_sheets: list[str] = field(default_factory=list)   # one pager only: sheets about another property
    include: list[str] = field(default_factory=list)   # file patterns that make sense for this row
    skipped: bool = False                   # analyst chose "None (skip this document)"
    saved: bool = False                     # re-applied from a saved choice
    reason: str = ""                        # why the automatic pick was made (debug log / review notes)
    auto_folder: Optional[str] = None
    auto_files: list[FileEntry] = field(default_factory=list)
    auto_sheets: list[str] = field(default_factory=list)

    @property
    def found(self) -> bool:
        return bool(self.files)

    @property
    def label(self) -> str:
        return f"{self.display_name} - {self.used_for}"

    @property
    def sections_text(self) -> str:
        return ", ".join(self.memo_sections)

    def files_text(self, limit: int = 3) -> str:
        if self.skipped:
            return "Skipped (you chose not to use this document)"
        if not self.files:
            return "Not found - double-click to choose"
        def show(e: FileEntry) -> str:
            parent = PurePosixPath(e.rel).parent
            return e.name if str(parent) == "." else f"{str(parent).replace('/', ' / ')} / {e.name}"
        if len(self.files) == 1:
            text = show(self.files[0])
        else:
            names = ", ".join(e.name for e in self.files[:limit])
            text = f"{len(self.files)} files: {names}" + (", ..." if len(self.files) > limit else "")
        if self.sheets:
            text += f"  [sheet{'s' if len(self.sheets) > 1 else ''}: {', '.join(self.sheets)}]"
        if self.saved:
            text += "  (saved choice)"
        elif not self.auto:
            text += "  (your choice)"
        return text

    def tooltip(self) -> str:
        return "\n".join(e.rel for e in self.files) + (f"\nSheets: {', '.join(self.sheets)}" if self.sheets else "")


@dataclass
class SourcesPlan:
    deal_path: str
    rows: dict[str, SourceRow]
    all_entries: list[FileEntry]
    storage: Any = None                     # the Storage that listed the folder (reused by generate)
    prior_deals: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)      # review notes found while checking (e.g. rejected sheets)

    # ------------------------------------------------------------------ for the dialog
    def folders(self) -> list[str]:
        dirs = sorted({e.rel for e in self.all_entries if e.is_dir}, key=str.lower)
        return [ROOT_LABEL] + dirs

    def files_in(self, folder: Optional[str], key: str, show_all: bool = False) -> list[FileEntry]:
        """Files directly inside `folder` (ROOT_LABEL/None/'.' = deal root), newest first."""
        rel = "" if folder in (None, "", ".", ROOT_LABEL) else folder
        row = self.rows.get(key)
        pats = [] if show_all or not row else row.include
        out = []
        for e in self.all_entries:
            if e.is_dir or e.parent_rel != (rel or "."):
                continue
            if e.name.startswith("~$"):
                continue
            if pats and not any(fnmatch.fnmatch(e.name.lower(), p.lower()) for p in pats):
                continue
            out.append(e)
        return sorted(out, key=lambda e: e.modified, reverse=True)

    # ------------------------------------------------------------------ overrides
    def set_files(self, key: str, files: list[FileEntry], sheets: list[str] | None = None) -> None:
        row = self.rows[key]
        row.files = list(files) if row.multi else list(files[:1])
        row.folder = row.files[0].parent_rel if row.files else row.folder
        row.sheets = list(sheets or [])
        row.auto, row.skipped, row.saved = False, False, False

    def skip(self, key: str) -> None:
        row = self.rows[key]
        row.files, row.sheets = [], []
        row.auto, row.skipped, row.saved = False, True, False

    def reset(self, key: str) -> None:
        row = self.rows[key]
        row.files, row.folder, row.sheets = list(row.auto_files), row.auto_folder, list(row.auto_sheets)
        row.auto, row.skipped, row.saved = True, False, False

    # ------------------------------------------------------------------ for the pipeline
    def selections(self) -> dict[str, Selection]:
        out = {}
        for k, r in self.rows.items():
            sel = Selection(k, r.display_name, r.folder if r.files else None, files=list(r.files),
                            note="skipped by the analyst" if r.skipped else r.reason, sheets=list(r.sheets),
                            exclude_sheets=[s for s in r.excluded_sheets if s not in r.sheets])
            out[k] = sel
        return out

    def summary_lines(self) -> list[str]:
        return [f"{r.display_name}: {r.files_text()}" for r in self.rows.values()]

    # ------------------------------------------------------------------ persistence
    def to_json(self) -> dict:
        rows = {}
        for k, r in self.rows.items():
            if r.auto:
                continue                    # only the analyst's own choices are remembered
            rows[k] = {"files": [e.rel for e in r.files], "sheets": r.sheets, "skipped": r.skipped}
        return {"deal_path": self.deal_path, "saved_at": datetime.now().isoformat(timespec="seconds"), "rows": rows}

    def save(self) -> Path:
        p = cache.plan_file(self.deal_path)
        data = self.to_json()
        if data["rows"]:
            p.write_text(json.dumps(data, indent=2), encoding="utf-8")
        elif p.exists():
            p.unlink()
        return p

    def apply_saved(self) -> list[str]:
        """Re-apply saved choices whose files still exist. Returns the keys re-applied."""
        p = cache.plan_file(self.deal_path)
        if not p.exists():
            return []
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            return []
        by_rel = {e.rel: e for e in self.all_entries if not e.is_dir}
        applied = []
        for k, saved in (data.get("rows") or {}).items():
            row = self.rows.get(k)
            if not row:
                continue
            if saved.get("skipped"):
                self.skip(k)
            else:
                files = [by_rel.get(r) for r in saved.get("files", [])]
                if not files or any(f is None for f in files):
                    continue                # a saved file was moved or deleted -> keep the automatic pick
                self.set_files(k, files, saved.get("sheets") or [])
            row.saved = True
            applied.append(k)
        return applied


# --------------------------------------------------------------------------- build
def build_plan(deal_path: str, selections: dict[str, Selection], entries: list[FileEntry], rules: dict,
               storage=None) -> SourcesPlan:
    src = (rules or {}).get("sources", {})
    rows: dict[str, SourceRow] = {}
    for key, sel in selections.items():
        rule = src.get(key, {})
        name, used = DEFAULT_NAMES.get(key, (sel.label, ""))
        why = ""
        if sel.files:
            c = next((c for c in sel.candidates if c.entry is sel.files[0]), None)
            why = "; ".join(c.reasons) if c else ""
        rows[key] = SourceRow(
            key=key, display_name=rule.get("display_name", name), used_for=rule.get("used_for", used),
            multi=rule.get("pick", "best") != "best", folder=sel.folder if sel.files else None,
            memo_sections=[str(s) for s in rule.get("memo_sections", []) or []],
            files=list(sel.files), include=list(rule.get("include", [])), reason=why or sel.note,
            auto_folder=sel.folder if sel.files else None, auto_files=list(sel.files))
    return SourcesPlan(deal_path, rows, entries, storage)


# --------------------------------------------------------------------------- one pager sheets
_SUFFIX = (r"st|street|ave|avenue|rd|road|blvd|boulevard|dr|drive|ln|lane|way|pl|place|ct|court|pkwy|parkway|"
           r"hwy|highway|ter|terrace|cir|circle|sq|square|trl|trail|tpke|turnpike|plz|plaza")
_DIR = r"(?:[NSEW]|north|south|east|west)\.?"
ADDRESS = re.compile(rf"\b(\d{{1,6}})\s+(?:{_DIR}\s+)?([A-Za-z][A-Za-z'.]*(?:\s+[A-Za-z][A-Za-z'.]*){{0,3}}?)\s+(?:{_SUFFIX})\b", re.I)
_TITLE_ADDR = re.compile(rf"^\s*(\d{{1,6}})\s+(?:{_DIR}\s+)?([A-Za-z][A-Za-z']+)", re.I)
_SKIP_WORDS = {"n", "s", "e", "w", "north", "south", "east", "west"}


def subject_street(deal_name: str) -> tuple[Optional[str], list[str]]:
    """'280 E Palmetto Park Rd' -> ('280', ['palmetto', 'park']); no leading number -> (None, [])."""
    m = re.match(r"\s*(\d{1,6})[A-Za-z]?\s+(.*)", deal_name or "")
    if not m:
        return None, []
    words = [w.lower() for w in re.findall(r"[A-Za-z]+", m.group(2))]
    core = [w for w in words if w not in _SKIP_WORDS and not re.fullmatch(_SUFFIX, w)]
    return m.group(1), core[:3]


def _mentions_subject(text: str, number: str, core: list[str]) -> bool:
    low = text.lower()
    return bool(re.search(rf"\b{re.escape(number)}\b", low)) and all(w in low for w in core[:2])


def other_addresses(text: str, number: Optional[str], core: list[str]) -> list[str]:
    """Street addresses in `text` that are not the subject property."""
    out = []
    for m in ADDRESS.finditer(text or ""):
        num, street = m.group(1), m.group(2).lower()
        if number and num == number and (not core or core[0] in street):
            continue
        out.append(m.group(0).strip())
    return out


def sheet_texts(path: Path, max_rows: int = 80, max_cols: int = 20) -> dict[str, str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        out = {}
        for ws in wb.worksheets:
            vals = []
            for row in ws.iter_rows(max_row=max_rows, max_col=max_cols, values_only=True):
                vals += [str(v) for v in row if v not in (None, "")]
            out[ws.title] = " | ".join(vals)
        return out
    finally:
        wb.close()


def pick_sheets(path: Path, deal_name: str) -> tuple[list[str], list[str], dict[str, str]]:
    """(all sheet names, sheets to use, {rejected sheet: reason}).

    Use the sheets that mention the subject's street number + street name; never a sheet that only
    mentions a different street address (another deal's one pager left in the workbook). Comps sheets
    are judged by title only, because comparables legitimately list other addresses.
    """
    texts = sheet_texts(path)
    names = list(texts)
    number, core = subject_street(deal_name)
    if not number:
        return names, names, {}
    matched, neutral, rejected = [], [], {}
    for title, text in texts.items():
        is_comps = bool(re.search(r"comp", title, re.I))
        if _mentions_subject(title + " " + text, number, core):
            matched.append(title)
            continue
        t = _TITLE_ADDR.match(title)
        if t and t.group(1) != number:
            rejected[title] = f"sheet is about {title.strip()}, not {deal_name}"
            continue
        others = [] if is_comps else other_addresses(text, number, core)
        if others:
            rejected[title] = f"mentions {others[0]}, not {deal_name}"
        else:
            neutral.append(title)
    comps = [t for t in neutral if re.search(r"comp", t, re.I)]
    chosen = (matched + comps) if matched else neutral
    return names, [n for n in names if n in chosen], rejected
