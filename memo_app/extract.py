"""Read each selected source file into a structured object (one LLM call per file)."""
from __future__ import annotations

import re
import traceback
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

from . import schemas as S
from .emails import EmailChain, chain_text
from .llm import LLM
from .progress import APPRAISAL, EMAILS, SPONSOR, TERM_SHEET, THIRD_PARTY
from .locator import Selection
from .storage import FileEntry, Storage
from .text_extract import DocText, last_pages_b64, load_doc, trim_to_relevant, xlsx_first_image

Log = Callable[[str], None]

COMMON = (
    "You extract data for a commercial real estate lender (Castellan Capital). "
    "Read the document carefully and fill every field you can find. Use null when a value is not stated - "
    "never guess. Money as plain US-dollar numbers (2240000), rates/percentages as percent numbers (9.64). "
)

INSTRUCTIONS = {
    "term_sheet": COMMON + (
        "This is a term sheet / letter of intent. 'Blended Loan Amount' is the total including any junior loan "
        "participant; capture the junior participant's name and amount separately. Read the Recourse / Guaranty "
        "clause to decide personal_guaranty / completion_guaranty / carveout_guaranty. The last page images, if "
        "supplied, show the signature block: set borrower_signed=true only if the borrower side is visibly signed. "
        "issuing_lender = the lender on the letterhead / 'Lender:' line. If that lender is not Castellan, or the "
        "term sheet describes a participation / pari passu / A-note split, Castellan only holds a share: fill "
        "structure, lead_lender, total_loan_amount and Castellan's share (castellan_share_amount / _pct) when stated."),
    "psa": COMMON + "This is a purchase and sale agreement. Find the purchase price (look for 'Purchase Price shall be').",
    "model": COMMON + (
        "This is Castellan's underwriting model exported cell by cell (R=row, letters=columns). The 'Model' sheet "
        "holds the Loan Overview (address, property facts), Sources & Uses, Loan Metrics and loan terms. 'Mezz' = "
        "junior loan, 'Rehab'/'CapEx' = capital expenditure budget. Infer property_type from the Loan Overview fields "
        "(e.g. loading docks/drive-in doors/clear height => industrial)."),
    "one_pager": COMMON + (
        "This is the memo 'one pager' workbook. Reproduce every table exactly as laid out, in reading order, with cell "
        "text formatted as displayed ($ with thousands separators, % signs, dates like 'December 11, 2025'). Typical "
        "tables: Loan Metrics, Loan Terms, Estimated Sources & Uses, Sales Comparables. For Sources & Uses use columns "
        "['Sources', '$', '%', 'Uses', '$']. Output each logical table ONCE: when the sheet shows the same table "
        "twice side by side (repeated column blocks), output it once with a single set of columns. Set each table's "
        "`sheet` to the sheet it is on (the '### Sheet:' line). Fill the single fields from ALL sheets: subject "
        "property_address, structure ('Castellan pari passu loan' -> pari_passu), lead_lender, total_loan_amount "
        "('Total Debt' / total senior loan), castellan_share_amount, loan_psf, comps_avg_psf, index_name, spread_pct, "
        "floor_rate_pct. Cell text is already formatted as Excel displays it - copy it as shown."),
    "pfs": COMMON + (
        "This is a personal financial statement (or several). Return one entry per individual. cash = cash and cash "
        "equivalents; marketable_securities = stocks/bonds/retirement; real_estate_assets = real estate owned value."),
    "background": COMMON + (
        "This is a background check (e.g. Eaglescreen) and/or credit report. Return one entry per person with the "
        "credit score (FICO/Vantage, mid score if several) and any adverse records. adverse_summary must be null when "
        "nothing adverse is reported."),
    "resume": COMMON + "This is a sponsor resume / bio. List the people it describes, their honorific if obvious, the company and any general contractor named.",
    "appraisal": COMMON + (
        "This is an appraisal report (possibly trimmed to the relevant pages). Capture the appraiser firm, all value "
        "conclusions (as is / as complete / as stabilized), building and site size, the sales comparables grid with "
        "adjusted $/SF and their average, and the neighbourhood / market / demographic facts with their numbers, "
        "periods and data sources (e.g. CoStar Q4 2025). For condominiums capture the gross retail sellout / "
        "aggregate retail value as gross_sellout_value. property_type must be one of the listed values."),
    "environmental": COMMON + (
        "This is an environmental report (Phase I ESA, Phase II ESA, HRDR / database review, or similar). "
        "status='completed' for a final report. recognized_environmental_conditions=true ONLY when the report "
        "recommends further investigation or remediation; a condition that is merely noted is not enough (describe "
        "it in conditions_identified). For a Phase II, no_further_assessment=true when it concludes that no further "
        "assessment or action is warranted."),
    "insurance": COMMON + "This is an insurance evidence / certificate / binder. Capture coverage limits.",
    "emails": COMMON + (
        "This is the email thread about a loan request between the broker / borrower and Castellan, oldest message "
        "first. Each message starts with a header line 'date | from | subject'. Quoted history has been removed, "
        "except inside forwards. Summarise what the thread says about the loan purpose, the business plan and the exit. "
        "IMPORTANT: amounts, leverage, rate and term that a broker or borrower REQUESTS are a request, not Castellan's "
        "final terms - put them only in requested_loan_amount / requested_ltv_or_ltc / requested_terms, and when the "
        "request changed during the thread use the latest one. Do not compute ratios. key_quotes must be verbatim "
        "(at most 25 words each) and end with ' - <sender name>, <YYYY-MM-DD>'. timeline entries start with the date."),
}

# Which progress step each source belongs to, and how to name it to the analyst.
GROUPS = {"term_sheet": TERM_SHEET, "psa": TERM_SHEET, "model": TERM_SHEET, "one_pager": TERM_SHEET,
          "appraisal": APPRAISAL, "pfs": SPONSOR, "background": SPONSOR, "resume": SPONSOR,
          "environmental": THIRD_PARTY, "insurance": THIRD_PARTY, "photos": THIRD_PARTY, "emails": EMAILS}
PLAIN_NAMES = {"term_sheet": "term sheet", "psa": "purchase agreement", "model": "underwriting model",
               "one_pager": "one pager", "appraisal": "appraisal", "pfs": "personal financial statements",
               "background": "background reports", "resume": "sponsor resume", "environmental": "environmental report",
               "insurance": "insurance documents", "photos": "property photos", "emails": "emails"}
IMPORTANT = {"term_sheet", "appraisal", "pfs", "background", "environmental", "insurance"}

APPRAISAL_KEYWORDS = [
    "neighborhood", "neighbourhood", "submarket", "market", "vacancy", "rent", "absorption", "demographic",
    "population", "household income", "median", "interstate", "highway", "access", "employer", "conclusion",
    "as is", "as complete", "as stabilized", "reconciliation", "final value", "comparable", "adjusted",
    "price per", "gross building", "rentable", "site", "location", "costar", "occupancy", "adr",
]


def dedupe_side_by_side(t: S.Table) -> S.Table:
    """A table copied twice next to itself (columns A B C A B C) -> one block of columns, unique rows."""
    cols = t.columns
    n = len(cols) // 2
    if n and len(cols) % 2 == 0 and [c.strip().lower() for c in cols[:n]] == [c.strip().lower() for c in cols[n:]]:
        rows, seen = [], set()
        for r in t.rows:
            for half in (r[:n], r[n:2 * n]):
                key = tuple(c.strip() for c in half)
                if any(key) and key not in seen:
                    seen.add(key)
                    rows.append(list(half))
        return S.Table(sheet=t.sheet, title=t.title, columns=list(cols[:n]), rows=rows, bold_rows=[])
    return t


_MONEY = re.compile(r"(\(?-?\$\s?\(?-?)([\d,]+\.\d{3,})")
_PERCENT = re.compile(r"(-?[\d,]*\.\d{3,})(\s?%)")


def tidy_table_numbers(t: S.Table) -> S.Table:
    """$ values with more than 2 decimals -> 0 decimals in a $/SF or PSF column, else 2; percentages -> max 2 decimals."""
    def fix(cell: str, header: str) -> str:
        per_sf = bool(re.search(r"\$/\s?SF|PSF|per\s+(sq|SF)", header or "", re.I))
        cell = _MONEY.sub(lambda m: m.group(1) + (f"{float(m.group(2).replace(',', '')):,.0f}" if per_sf
                                                   else f"{float(m.group(2).replace(',', '')):,.2f}"), cell)
        return _PERCENT.sub(lambda m: f"{float(m.group(1).replace(',', '')):.2f}".rstrip("0").rstrip(".") + m.group(2), cell)
    rows = []
    for r in t.rows:
        label = r[0] if r else ""
        rows.append([fix(str(c), t.columns[i] if i < len(t.columns) else label) for i, c in enumerate(r)])
    return t.model_copy(update={"rows": rows})


@dataclass
class Extracted:
    term_sheet: Optional[S.TermSheet] = None
    psa: Optional[S.PSA] = None
    model: Optional[S.UWModel] = None
    one_pager: Optional[S.OnePager] = None
    pfs: Optional[S.PFSResult] = None
    background: Optional[S.BackgroundResult] = None
    resume: Optional[S.ResumeFacts] = None
    appraisal: Optional[S.Appraisal] = None
    environmental: Optional[S.Environmental] = None            # the Phase I (or the only report)
    environmental_reports: list[S.Environmental] = field(default_factory=list)   # every report read (Phase I + II)
    insurance: Optional[S.Insurance] = None
    adverse_search: Optional[S.AdverseSearch] = None
    email_insights: Optional[S.EmailInsights] = None
    emails: Optional[EmailChain] = None
    resume_text: str = ""
    appraisal_text: str = ""
    plan_notes: list[str] = field(default_factory=list)   # review notes from the sources check
    error_items: list[tuple[str, str]] = field(default_factory=list)   # (source key, error) for the review list
    photos: list[tuple[str, bytes]] = field(default_factory=list)
    provenance: dict[str, list[str]] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)


class Extractor:
    def __init__(self, storage: Storage, llm: LLM, settings, log: Log = print):
        self.storage = storage
        self.llm = llm
        self.s = settings
        self.log = log
        self.doc_cfg = settings.get("documents", {}) or {}
        self.max_chars = int(self.doc_cfg.get("max_chars_per_document", 180000))
        self.failed: dict[str, list[str]] = {}       # progress step -> plain names of documents that failed
        self.missing: dict[str, list[str]] = {}      # progress step -> important documents not found

    # ------------------------------------------------------------------ utils
    def _doc(self, e: FileEntry, sheets: list[str] | None = None) -> tuple[Path, DocText]:
        p = self.storage.fetch(e)
        d = load_doc(p, e.rel, self.doc_cfg, sheets=sheets)
        if d.note:
            self.log(f"  note: {e.rel}: {d.note}")
        return p, d

    def _payload(self, d: DocText, keywords: list[str] | None = None) -> tuple[str, list[str]]:
        text = trim_to_relevant(d, keywords, self.max_chars) if keywords else d.text[: self.max_chars]
        return f"FILE: {d.name}\n\n{text}", d.images_b64

    def _extract_file(self, key, schema, e: FileEntry, extra_images=None, keywords=None):
        _, d = self._doc(e)
        text, imgs = self._payload(d, keywords)
        return self.llm.extract(schema, INSTRUCTIONS[key], text, (imgs or []) + (extra_images or []))

    # ------------------------------------------------------------------ term sheet
    def _term_sheet(self, sel: Selection, out: Extracted):
        e = sel.files[0]
        p = self.storage.fetch(e)
        imgs = last_pages_b64(p, 2) if e.suffix == ".pdf" else []
        out.term_sheet = self._extract_file("term_sheet", S.TermSheet, e, extra_images=imgs)

    # ------------------------------------------------------------------ multi-file merges
    def _pfs(self, sel: Selection, out: Extracted):
        people: dict[str, S.SponsorPFS] = {}
        for e in sel.files:
            res: S.PFSResult = self._extract_file("pfs", S.PFSResult, e)
            for sp in res.sponsors:
                k = sp.full_name.strip().lower()
                if k not in people or (sp.net_worth or 0) > (people[k].net_worth or 0):
                    people[k] = sp
        out.pfs = S.PFSResult(sponsors=list(people.values()))

    def _background(self, sel: Selection, out: Extracted):
        merged: dict[str, S.PersonBackground] = {}
        provider = None
        for e in sel.files:
            res: S.BackgroundResult = self._extract_file("background", S.BackgroundResult, e)
            provider = provider or res.provider
            for p in res.people:
                k = p.full_name.strip().lower()
                if k in merged:
                    old = merged[k]
                    data = old.model_dump()
                    for f, v in p.model_dump().items():
                        if data.get(f) in (None, False, "") and v not in (None, ""):
                            data[f] = v
                    merged[k] = S.PersonBackground(**data)
                else:
                    merged[k] = p
        out.background = S.BackgroundResult(provider=provider, people=list(merged.values()))

    def _insurance(self, sel: Selection, out: Extracted):
        best: dict = {}
        for e in sel.files:
            ins: S.Insurance = self._extract_file("insurance", S.Insurance, e)
            for f, v in ins.model_dump().items():
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    best[f] = max(best.get(f) or 0, v)
                elif f == "is_bound":
                    best[f] = bool(best.get(f)) or bool(v)
                elif best.get(f) is None:
                    best[f] = v
        out.insurance = S.Insurance(**{k: best.get(k) for k in S.Insurance.model_fields})

    def _resume(self, sel: Selection, out: Extracted):
        texts = []
        for e in sel.files:
            _, d = self._doc(e)
            texts.append(f"FILE: {e.rel}\n{d.text[:60000]}")
        out.resume_text = "\n\n".join(texts)
        if out.resume_text.strip():
            out.resume = self.llm.extract(S.ResumeFacts, INSTRUCTIONS["resume"], out.resume_text[:120000])

    def _appraisal(self, sel: Selection, out: Extracted):
        _, d = self._doc(sel.files[0])
        out.appraisal_text = trim_to_relevant(d, APPRAISAL_KEYWORDS, self.max_chars)
        out.appraisal = self.llm.extract(S.Appraisal, INSTRUCTIONS["appraisal"],
                                         f"FILE: {d.name}\n\n{out.appraisal_text}", d.images_b64[:6])

    def _one_pager(self, sel: Selection, out: Extracted):
        """Facts (share, total debt, lead lender, rate terms, loan metrics) from EVERY sheet except those about another
        property; only the tables of the sheets chosen in the Sources table go in the memo."""
        p = self.storage.fetch(sel.files[0])
        keep = None
        if sel.exclude_sheets and p.suffix.lower() in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True)
            keep = [n for n in wb.sheetnames if n not in sel.exclude_sheets]
            wb.close()
        _, d = self._doc(sel.files[0], sheets=keep)
        text, imgs = self._payload(d)
        op = self.llm.extract(S.OnePager, INSTRUCTIONS["one_pager"], text, imgs)
        tables = [tidy_table_numbers(dedupe_side_by_side(t)) for t in op.tables]
        if sel.sheets:
            tables = [t for t in tables if not t.sheet or t.sheet in sel.sheets]
        out.one_pager = op.model_copy(update={"tables": tables})

    def _environmental(self, sel: Selection, out: Extracted):
        """Every selected report (Phase I + Phase II); the Phase I is the main one."""
        reports = [self._extract_file("environmental", S.Environmental, e) for e in sel.files]
        out.environmental_reports = reports
        out.environmental = next((r for r in reports if r.phase == "phase_i"), reports[0] if reports else None)

    def _emails(self, chain: EmailChain, out: Extracted):
        text = chain_text(chain)
        if text.strip():
            out.email_insights = self.llm.extract(S.EmailInsights, INSTRUCTIONS["emails"], text)

    def _photos(self, sel: Selection, out: Extracted):
        for e in sel.files:
            out.photos.append((e.name, self.storage.fetch(e).read_bytes()))

    # ------------------------------------------------------------------ run all
    def run(self, selections: dict[str, Selection], emails: Optional[EmailChain] = None, rep=None,
            defer: set[str] | None = None) -> Extracted:
        """Read every selected source in parallel. With `rep`, each progress step (term sheet, appraisal,
        sponsor documents, third-party reports, emails) is marked done as soon as its documents are read;
        steps in `defer` are left open for the caller (it calls group_status() and done() itself)."""
        out = Extracted(emails=emails)
        simple = {"psa": S.PSA, "model": S.UWModel}
        tasks: dict[str, Callable] = {}
        defer = defer or set()

        for key, sel in selections.items():
            if not sel.found:
                if key in IMPORTANT and "skipped" not in (sel.note or ""):
                    self.missing.setdefault(GROUPS.get(key, ""), []).append(PLAIN_NAMES.get(key, key))
                continue
            out.provenance[key] = [f.rel for f in sel.files]
            if key == "term_sheet":
                tasks[key] = lambda sel=sel: self._term_sheet(sel, out)
            elif key in simple:
                def job(key=key, sel=sel):
                    setattr(out, key, self._extract_file(key, simple[key], sel.files[0]))
                tasks[key] = job
            elif key == "one_pager":
                tasks[key] = lambda sel=sel: self._one_pager(sel, out)
            elif key == "environmental":
                tasks[key] = lambda sel=sel: self._environmental(sel, out)
            elif key == "pfs":
                tasks[key] = lambda sel=sel: self._pfs(sel, out)
            elif key == "background":
                tasks[key] = lambda sel=sel: self._background(sel, out)
            elif key == "insurance":
                tasks[key] = lambda sel=sel: self._insurance(sel, out)
            elif key == "resume":
                tasks[key] = lambda sel=sel: self._resume(sel, out)
            elif key == "appraisal":
                tasks[key] = lambda sel=sel: self._appraisal(sel, out)
            elif key == "photos":
                tasks[key] = lambda sel=sel: self._photos(sel, out)
        if emails and emails.messages:          # local .eml/.msg files and/or the Outlook thread
            tasks["emails"] = lambda: self._emails(emails, out)

        pending: dict[str, int] = {}
        for key in tasks:
            pending[GROUPS.get(key, "")] = pending.get(GROUPS.get(key, ""), 0) + 1
        if rep:
            for group in (TERM_SHEET, APPRAISAL, SPONSOR, THIRD_PARTY):
                rep.step(group)
                if not pending.get(group) and group not in defer:
                    ok, note = self.group_status(group)
                    rep.done(group, ok=ok, note=note)

        workers = int(self.s.get("openai.max_parallel_requests", 4))
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            futs = {pool.submit(fn): key for key, fn in tasks.items()}
            for fut in as_completed(futs):
                key = futs[fut]
                label = selections[key].label if key in selections else key
                group = GROUPS.get(key, "")
                try:
                    fut.result()
                    self.log(f"  read {label}")
                except Exception as exc:
                    msg = f"{label}: {exc}"
                    out.errors.append(msg)
                    out.error_items.append((key, str(exc)))
                    self.failed.setdefault(group, []).append(PLAIN_NAMES.get(key, key))
                    self.log(f"  ERROR {msg}")
                    self.log(traceback.format_exc())
                pending[group] -= 1
                if rep and pending[group] == 0 and group not in defer:
                    ok, note = self.group_status(group)
                    rep.done(group, ok=ok, note=note)

        # property photo embedded in the model workbook, used when Site Visit has none
        if not out.photos and selections.get("model") and selections["model"].found:
            try:
                img = xlsx_first_image(self.storage.fetch(selections["model"].files[0]))
                if img:
                    out.photos.append(("model photo", img))
            except Exception as exc:
                self.log(f"  model photo not read: {exc}")
        return out

    def group_status(self, group: str) -> tuple[bool, str]:
        """(ok, plain note) for a progress step, e.g. (False, "Couldn't read the appraisal - ...")."""
        failed, missing = self.failed.get(group, []), self.missing.get(group, [])
        if failed:
            return False, f"Couldn't read the {' and '.join(failed)} - the memo will mark that section for review."
        if missing:
            return False, f"No {' or '.join(missing)} found - the memo will mark that section for review."
        return True, ""

    def adverse_search(self, names: list[str], city_state: str, out: Extracted) -> None:
        if not names:
            return
        prompt = (
            "Search the web for adverse information about these real estate sponsors: "
            f"{'; '.join(names)} (active around {city_state or 'the USA'}). Look for fraud, bankruptcy, foreclosure, "
            "lawsuits, liens, criminal charges, regulatory actions or negative press. Be careful about people with "
            "the same name - only count results that clearly match a real estate investor/developer. Return one "
            "finding per item (person, one-sentence finding, category, severity, source URL and page title); "
            "neutral mentions get severity 'info'. adverse_found=false if nothing credible and adverse is found."
        )
        try:
            out.adverse_search = self.llm.web_search(S.AdverseSearch, prompt)
        except Exception as exc:
            out.errors.append(f"Online adverse search failed: {exc}")
