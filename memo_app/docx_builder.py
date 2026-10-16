"""Assemble the memo .docx on top of resources/templates/memo_base.docx
(Castellan logo, fonts, header date, page-number footer copied from the team's template)."""
from __future__ import annotations

import io
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import paths
from .facts import GEN, HL, MAP, PICS, R_CON, R_ENV, R_INS, R_IR, R_LTV, R_SP, SP_BG, SP_FIN, TO
from .render import AI_C, AI_O, FLAG_C, FLAG_O, UID_TOKEN, VAR_C, VAR_O, render, when

NAVY = "002060"
RED = RGBColor(0xC0, 0x00, 0x00)
BODY_PT = 11
TOKEN = re.compile(r"(\*\*|" + "|".join(map(re.escape, (VAR_O, VAR_C, FLAG_O, FLAG_C, AI_O, AI_C))) + ")")


RISK_SECTIONS = {"ltv": R_LTV, "sponsor": R_SP, "construction": R_CON, "environmental": R_ENV,
                 "interest_rate": R_IR, "insurance": R_INS}


class MemoDoc:
    """Only red remarks are highlighted (red highlight, white bold). Tool-filled values (⟦⟧) and
    AI text (⟬⟭) are tracked internally but print as normal text.

    Each red remark is printed as "[LEVEL Rn: title]": the flag it points to gets the next ID the first time
    it appears, so IDs follow document order and match the review workbook (same Flag objects)."""

    def __init__(self, facts=None):
        self.doc = docx.Document(str(paths.resource("templates", "memo_base.docx")))
        self.facts = facts
        self.section = GEN
        self.next_id = 1

    # ------------------------------------------------------------------ red remarks with IDs
    def _number(self, fl) -> None:
        if not fl.id:
            fl.id = f"R{self.next_id}"
            self.next_id += 1
        fl.inline = True

    def _remark(self, tok: str) -> str:
        if self.facts is None:
            return tok
        m = UID_TOKEN.fullmatch(tok.strip())
        if m:
            fl = next((x for x in self.facts.flags if x.uid == int(m.group(1))), None)
            if fl is None:
                return ""
        else:      # a free-text remark: make it a flag so it has an ID and a workbook row
            text = tok.strip().strip("[]").strip()
            if not text:
                return tok
            level = "missing" if text.upper().startswith("MISSING") else ("check" if text.upper().startswith("CHECK") else "review")
            title = text.split(":", 1)[1].strip() if ":" in text[:9] else text
            fl = self.facts.flag(level, title, title, self.section)
        self._number(fl)
        return fl.remark_text()

    def number_remaining(self) -> None:
        """Flags that never appear in the body get the IDs after the inline ones."""
        for fl in (self.facts.flags if self.facts is not None else []):
            if not fl.id:
                fl.id = f"R{self.next_id}"
                self.next_id += 1

    # ------------------------------------------------------------------ text with markers
    def add_rich(self, p, text: str, size: int = BODY_PT, bold: bool = False, italic: bool = False,
                 red: bool = False):
        state = {"b": False, "var": 0, "flag": 0, "ai": 0}
        for tok in TOKEN.split(text or ""):
            if not tok:
                continue
            if tok == "**":
                state["b"] = not state["b"]; continue
            if tok in (VAR_O, VAR_C):
                state["var"] += 1 if tok == VAR_O else -1; continue
            if tok in (FLAG_O, FLAG_C):
                state["flag"] += 1 if tok == FLAG_O else -1; continue
            if tok in (AI_O, AI_C):
                state["ai"] += 1 if tok == AI_O else -1; continue
            if state["flag"] > 0:
                tok = self._remark(tok)
                if not tok:
                    continue
            r = p.add_run(tok)
            r.font.size = Pt(size)
            r.bold = bold or state["b"] or state["flag"] > 0
            r.italic = italic
            if state["flag"] > 0:
                r.font.highlight_color = WD_COLOR_INDEX.RED
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif red:
                r.font.color.rgb = RED
        return p

    # ------------------------------------------------------------------ paragraphs
    @staticmethod
    def _spacing(p, before=0, after=6, justify=True):
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        if justify:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def set_date(self, date_text: str):
        for s in self.doc.sections:
            for part in (s.header, s.first_page_header):
                for p in part.paragraphs:
                    for r in p.runs:
                        if "{{MEMO_DATE}}" in r.text:
                            r.text = r.text.replace("{{MEMO_DATE}}", date_text)

    def title(self, title: str, address: str):
        for text in (title, address):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            self.add_rich(p, text, size=12, bold=True)
        self._spacing(self.doc.add_paragraph(), after=2, justify=False)

    def heading(self, text: str):
        p = self.doc.add_paragraph()
        self._spacing(p, before=8, after=2)
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "auto")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)
        self.add_rich(p, text, bold=True)
        return p

    def para(self, text: str, italic=False):
        for chunk in [c for c in re.split(r"\n\s*\n", text or "") if c.strip()]:
            p = self.doc.add_paragraph()
            self._spacing(p)
            self.add_rich(p, chunk.strip(), italic=italic)

    def bullet(self, text: str, label: str | None = None):
        try:
            p = self.doc.add_paragraph(style="List Paragraph")
        except KeyError:
            p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl, numId = OxmlElement("w:ilvl"), OxmlElement("w:numId")
        ilvl.set(qn("w:val"), "0"); numId.set(qn("w:val"), "1")   # numId 1 = bullet list in the base template
        numPr.append(ilvl); numPr.append(numId)
        pPr.append(numPr)
        self._spacing(p, after=4)
        if label:
            self.add_rich(p, label + " ", bold=True)
        self.add_rich(p, text)

    # ------------------------------------------------------------------ tables
    @staticmethod
    def _shade(cell, hex_fill: str):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    @staticmethod
    def _borders(table):
        tblPr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:color"), "000000")
            b.append(el)
        tblPr.append(b)

    def table(self, t):
        cols = max([len(t.columns)] + [len(r) for r in t.rows] + [1])
        n_rows = 1 + (1 if t.columns else 0) + len(t.rows)
        table = self.doc.add_table(rows=n_rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._borders(table)
        table.autofit = False
        if cols == 2:
            col_w = [Inches(2.2), Inches(2.6)]
        else:  # share 6.5" in proportion to the longest text in each column
            longest = [max([len(str(r[i])) if i < len(r) else 0 for r in t.rows] +
                           [len(t.columns[i]) if i < len(t.columns) else 0]) for i in range(cols)]
            weights = [min(max(n + 3, 10), 32) for n in longest]
            col_w = [Inches(6.5 * w / sum(weights)) for w in weights]
        for i, col in enumerate(table.columns):
            col.width = col_w[i]
        for row in table.rows:
            for i, c in enumerate(row.cells):
                c.width = col_w[i]
        title = table.rows[0].cells[0].merge(table.rows[0].cells[-1])
        self._shade(title, NAVY)
        p = title.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t.title)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(10), RGBColor(0xFF, 0xFF, 0xFF)

        start = 1
        if t.columns:
            for i, h in enumerate(t.columns):
                cp = table.rows[1].cells[i].paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_after = Pt(0)
                rr = cp.add_run(h); rr.bold, rr.underline, rr.font.size = True, True, Pt(9.5)
            start = 2
        for ri, row in enumerate(t.rows):
            for ci in range(cols):
                val = row[ci] if ci < len(row) else ""
                cp = table.rows[start + ri].cells[ci].paragraphs[0]
                cp.paragraph_format.space_after = Pt(0)
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 or (t.columns and ci == 3) else WD_ALIGN_PARAGRAPH.CENTER
                rr = cp.add_run(str(val)); rr.font.size = Pt(9.5)
                rr.bold = ri in (t.bold_rows or [])
        self._spacing(self.doc.add_paragraph(), after=4, justify=False)

    # ------------------------------------------------------------------ pictures
    def pictures(self, photos: list[tuple[str, bytes]]):
        ok = []
        for name, data in photos:
            if data[:3] == b"\xff\xd8\xff" or data[:8] == b"\x89PNG\r\n\x1a\n":
                ok.append((name, data))
        if not ok:
            mark = self.facts.new_remark("missing", "insert property photos", "No property photos were found (Site "
                                         "Visit folder or the model) - insert photos of the property.", PICS, "photos",
                                         key="photos_missing") if self.facts is not None else f"{FLAG_O}[Insert property photos]{FLAG_C}"
            self.para(mark)
            return
        for i in range(0, len(ok), 2):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for _, data in ok[i:i + 2]:
                p.add_run().add_picture(io.BytesIO(data), width=Inches(3.1))
                p.add_run("  ")

    def save(self, path: str):
        self.doc.save(path)


# --------------------------------------------------------------------------- assemble
def build_memo(path_out: str, f, sections: dict[str, str], tables, photos, bp: dict):
    """Build the .docx. Assigns the review IDs (R1, R2, ...) to f.flags in document order - build the review
    workbook AFTER this so both use the same IDs."""
    m = MemoDoc(f)
    H = bp["headings"]
    m.set_date(f.memo_date)
    m.section = TO
    m.title(bp.get("title", "LOAN INVESTMENT MEMORANDUM"),
            f"{VAR_O}{f.address}{VAR_C}" if f.address else f.remark("address_missing", "missing", "property address", TO))

    m.heading(H["transaction_overview"])
    m.para(sections.get("transaction_overview", ""))

    m.heading(H["sponsorship_overview"])
    sp = bp["sponsorship"]
    m.section = "Sponsorship - Business background"
    m.bullet(sections.get("business_background", ""), sp["business_background_label"])
    m.section = SP_FIN
    m.bullet(render(sp["personal_financials"], f, SP_FIN), sp["personal_financials_label"])
    m.section = SP_BG
    m.bullet(render(sp["background"], f, SP_BG), sp["background_label"])

    m.heading(H["submarket_overview"])
    m.section = "Submarket Overview"
    m.para(sections.get("submarket_overview", ""))

    m.heading(H["business_plan"])
    m.section = "Business Plan"
    m.para(sections.get("business_plan", ""))

    m.heading(H["risks"])
    for risk in bp["risks"]:
        if when(risk.get("when", "true"), f):
            m.section = RISK_SECTIONS.get(risk.get("key"), "Risks")
            text = render(risk["text"], f, m.section)
            if text:
                m.bullet(text, risk["label"])

    m.heading(H["loan_financials"])
    m.section = "Loan Financials"
    m.para(bp["loan_financials_intro"], italic=True)
    for t in tables:
        m.table(t)

    m.heading(H["investment_highlights"])
    m.section = HL
    m.para(bp["highlights_intro"], italic=True)
    for h in bp["highlights"]:
        if when(h.get("when", "true"), f):
            m.bullet(render(h["text"], f, HL))

    m.heading(H["pictures"])
    m.section = PICS
    m.pictures(photos)
    m.heading(H["map"])
    m.section = MAP
    m.para(f.new_remark("missing", "insert map", "Insert a map of the property location.", MAP, key="map_missing"))
    m.number_remaining()
    m.save(path_out)
    return path_out
