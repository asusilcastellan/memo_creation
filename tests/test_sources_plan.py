"""Sources table (memo_app/sources_plan.py): overrides, saved choices, one-pager sheets, no rescan."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook  # noqa: E402

from memo_app import schemas as S  # noqa: E402
from memo_app.facts import UserInputs  # noqa: E402
from memo_app.settings import load_settings  # noqa: E402
from memo_app.sources_plan import pick_sheets, subject_street  # noqa: E402
from memo_app.storage import LocalStorage  # noqa: E402
from test_offline import CANNED, StubLLM, make_deal, offline_rates  # noqa: E402


def make_one_pager(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "150 E Aurora"
    ws["A1"], ws["A2"], ws["B2"] = "150 E Aurora St, Waterbury, CT", "Loan Amount:", 2240000
    other = wb.create_sheet("77 Mott")                          # another deal's one pager left in the workbook
    other["A1"], other["A2"] = "77 Mott Street, New York, NY", "Loan Amount: 9,000,000"
    comps = wb.create_sheet("Comps")
    comps["A1"] = "1050 Buckingham Street, Watertown, CT"
    misc = wb.create_sheet("Notes")
    misc["A1"] = "Prepared by 12 Main Street Partners"         # mentions another address only
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_subject_street_and_sheet_filtering(tmp_path):
    assert subject_street("280 E Palmetto Park Rd") == ("280", ["palmetto", "park"])
    assert subject_street("Some Portfolio") == (None, [])
    wb = make_one_pager(tmp_path / "one pager.xlsx")
    names, chosen, rejected = pick_sheets(wb, "150 E Aurora St")
    assert names == ["150 E Aurora", "77 Mott", "Comps", "Notes"]
    assert chosen == ["150 E Aurora", "Comps"]
    assert set(rejected) == {"77 Mott", "Notes"} and "not 150 E Aurora St" in rejected["77 Mott"]


def test_override_model_and_sheet_generate_without_rescan(tmp_path, monkeypatch):
    import memo_app.pipeline as pipeline
    deal = make_deal(tmp_path)
    make_one_pager(deal / "Memo" / "150 E Aurora one pager.xlsx")
    walks = []
    orig_walk = LocalStorage.walk
    monkeypatch.setattr(LocalStorage, "walk", lambda self, root: (walks.append(root), orig_walk(self, root))[1])
    monkeypatch.setattr(pipeline, "LLM", StubLLM)
    monkeypatch.setitem(CANNED, "OnePager", S.OnePager.model_validate(_one_pager_payload()))

    settings = load_settings()
    settings.openai_api_key = "stub"
    ui = UserInputs(deal_path=str(deal), loan_type="acquisition", output_dir=str(tmp_path / "out"))
    plan = pipeline.scan(ui, settings)
    assert plan.rows["one_pager"].sheets == ["150 E Aurora", "Comps"]            # 77 Mott dropped automatically
    assert any("77 Mott" in n for n in plan.notes)

    by_name = {e.name: e for e in plan.all_entries}
    plan.set_files("model", [by_name["150 E Aurora St - Net Funding.xlsx"]])
    plan.set_files("one_pager", [by_name["150 E Aurora one pager.xlsx"]], sheets=["150 E Aurora"])
    plan.save()
    with offline_rates():
        out = pipeline.generate(ui, settings, plan)
    assert out.memo.exists() and len(walks) == 1                                 # generate did not rescan
    sent = StubLLM.last.extract_inputs
    assert "150 E Aurora St - Net Funding.xlsx" in sent["UWModel"]
    assert "Sheet: 150 E Aurora" in sent["OnePager"] and "77 Mott" not in sent["OnePager"]

    again = pipeline.scan(ui, settings)                                           # saved choice re-applied
    assert again.rows["model"].files[0].name == "150 E Aurora St - Net Funding.xlsx" and again.rows["model"].saved
    assert "(saved choice)" in again.rows["model"].files_text()
    again.reset("model")
    assert again.rows["model"].files[0].name == "150 E Aurora St Model.xlsx" and again.rows["model"].auto
    again.skip("photos")
    assert again.selections()["photos"].files == [] and "Skipped" in again.rows["photos"].files_text()


def _one_pager_payload() -> dict:
    """A OnePager with every field null except the tables (fields are added over time)."""
    data = {k: None for k in S.OnePager.model_fields}
    data["tables"] = [S.Table(sheet=None, title="Loan Metrics", columns=[], rows=[["Loan Amount:", "$2,240,000"]], bold_rows=[])]
    return data


def test_rows_carry_memo_sections_from_sources_yaml(tmp_path):
    import memo_app.pipeline as pipeline
    from memo_app.settings import load_yaml_resource
    deal = make_deal(tmp_path)
    plan = pipeline.scan(UserInputs(deal_path=str(deal), loan_type="acquisition"), load_settings())
    rules = load_yaml_resource("config", "sources.yaml")["sources"]
    for key, row in plan.rows.items():
        assert row.memo_sections == rules[key]["memo_sections"] and row.memo_sections, key
    assert plan.rows["term_sheet"].sections_text == \
        "Transaction Overview, Business Plan, Risks - Sponsor / Construction / Interest Rate, Loan Financials"
    assert plan.rows["photos"].memo_sections == ["Pictures"]
