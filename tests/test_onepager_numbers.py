"""Part C fixes from the 280 E Palmetto draft: one-pager facts from all sheets, participations without a known
share, and Excel numbers as Excel displays them. Offline."""
from __future__ import annotations

import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook  # noqa: E402

from memo_app import schemas as S  # noqa: E402
from memo_app.extract import Extracted, Extractor, tidy_table_numbers  # noqa: E402
from memo_app.facts import UserInputs, build_facts  # noqa: E402
from memo_app.locator import Selection  # noqa: E402
from memo_app.render import render, strip_markers, when  # noqa: E402
from memo_app.settings import Settings, load_yaml_resource  # noqa: E402
from memo_app.storage import FileEntry, LocalStorage  # noqa: E402
from memo_app.text_extract import format_cell, xlsx_to_doc  # noqa: E402
from memo_app.writer import Writer  # noqa: E402
from test_accuracy import _x  # noqa: E402
from test_language import ScriptedLLM  # noqa: E402

BP = load_yaml_resource("templates", "boilerplate.yaml")
SETTINGS = Settings(raw={})
MAXIM_TS = {"issuing_lender": "Maxim Capital", "structure": "pari_passu", "lead_lender": "Maxim Capital",
            "total_loan_amount": 70_000_000, "castellan_share_amount": None, "castellan_share_pct": None,
            "blended_loan_amount": 70_000_000, "senior_loan_amount": None, "junior_loan_amount": None,
            "junior_lender_name": None}


def _one_pager_workbook(path: Path) -> Path:
    wb = Workbook()
    terms = wb.active
    terms.title = "Loan Terms"
    terms["A1"], terms["A2"], terms["B2"] = "280 E Palmetto Park Rd", "Castellan pari passu loan", 25_000_000
    terms["B2"].number_format = "$#,##0"
    terms["A3"], terms["B3"] = "Total Debt", 70_000_000
    comps = wb.create_sheet("Sales Comps")
    comps["A1"], comps["B1"], comps["C1"] = "Address", "$/SF", "Adj. $/SF"
    comps["A2"], comps["B2"] = "100 Main Street", 869.065868
    comps["B2"].number_format = "$#,##0"
    other = wb.create_sheet("77 Mott")
    other["A1"] = "77 Mott Street - Castellan loan $9,000,000"
    wb.save(path)
    return path


class OnePagerLLM:
    def __init__(self):
        self.inputs = {}

    def extract(self, schema, instructions, text=None, images_b64=None):
        self.inputs[schema.__name__] = text
        return S.OnePager(
            tables=[S.Table(sheet="Loan Terms", title="Loan Terms", columns=[], rows=[["Loan Amount", "$70,000,000"]], bold_rows=[]),
                    S.Table(sheet="Sales Comps", title="Sales Comparables", columns=["Address", "$/SF"],
                            rows=[["100 Main Street", "$869.065868"]], bold_rows=[])],
            property_address="280 E Palmetto Park Rd", structure="pari_passu", lead_lender="Maxim Capital",
            total_loan_amount=70_000_000, castellan_share_amount=25_000_000, castellan_share_pct=None, loan_psf=None,
            comps_avg_psf=None, index_name=None, spread_pct=None, floor_rate_pct=None)


# --------------------------------------------------------------------------- C1
def test_share_found_on_an_unselected_sheet(tmp_path):
    wb = _one_pager_workbook(tmp_path / "280 Palmetto one pager.xlsx")
    entry = FileEntry(str(wb), wb.name, wb.name, False, datetime(2025, 1, 1, tzinfo=timezone.utc))
    sel = Selection("one_pager", "One pager", ".", files=[entry], sheets=["Sales Comps"], exclude_sheets=["77 Mott"])
    llm = OnePagerLLM()
    x = Extractor(LocalStorage(), llm, SETTINGS, log=lambda m: None).run({"one_pager": sel})
    sent = llm.inputs["OnePager"]
    assert "Sheet: Loan Terms" in sent and "Sheet: Sales Comps" in sent and "77 Mott" not in sent
    assert "B=$25,000,000" in sent and "B=$869" in sent                      # numbers as Excel shows them
    assert [t.title for t in x.one_pager.tables] == ["Sales Comparables"]   # only the selected sheet's tables
    assert x.one_pager.tables[0].rows[0][1] == "$869"                        # $/SF column -> 0 decimals
    full = _x(term_sheet=MAXIM_TS, appraisal={"as_is_value": 100_000_000, "as_complete_value": None})
    full.one_pager = x.one_pager
    f = build_facts(full, UserInputs(deal_path=".", loan_type="acquisition"), SETTINGS)
    assert f.loan == 25_000_000 and f.total_loan == 70_000_000 and not f.find("share_missing")


def test_no_share_anywhere_uses_will_participate_and_missing_flag():
    x = _x(term_sheet=MAXIM_TS, model={"senior_loan": 70_000_000, "junior_loan": None},
           appraisal={"as_is_value": 100_000_000, "as_complete_value": None})
    f = build_facts(x, UserInputs(deal_path=".", loan_type="acquisition"), SETTINGS)
    assert f.is_participation and f.loan is None and round(f.memo_ltv, 2) == 0.70
    share = f.find("share_missing")
    assert share and share[0].level == "missing" and share[0].title == "Castellan's share"
    assert not f.find("loan_missing")
    opener = strip_markers(render(BP["openers"]["structure"]["participation_no_share"], f), f)
    assert opener == ("Castellan Capital will participate in a $70,000,000 senior loan led by Maxim Capital "
                      "[MISSING: Castellan's share]")
    risk = strip_markers(render(next(r["text"] for r in BP["risks"] if r["key"] == "ltv"), f), f)
    assert "in which Castellan participates, represents an estimated Loan-to-Value ratio of 70%" in risk
    assert "holds a" not in risk and "$70,000,000 pari passu" not in risk
    hl = " ".join(strip_markers(render(h["text"], f), f) for h in BP["highlights"] if when(h.get("when", "true"), f))
    assert "Castellan participates in a $70,000,000 senior loan" in hl and "pari passu position" not in hl
    # the writer picks the variant and keeps the red remark in front of the model's continuation
    plain = "Castellan Capital will participate in a $70,000,000 senior loan led by Maxim Capital"
    llm = ScriptedLLM([plain + " to finance the acquisition of the property located at 150 E Aurora St, Waterbury, CT. "
                       + "The property is an industrial building of 89,855 square feet. " * 4])
    w = Writer(llm, Settings(raw={"memo": {"few_shot_examples_per_section": 1}}), BP, log=lambda m: None)
    out = strip_markers(w.write_section("transaction_overview", f, x), f)
    assert out.startswith(plain + " [MISSING: Castellan's share] to finance the acquisition")
    assert "a $70,000,000 participation in a $70,000,000" not in out


# --------------------------------------------------------------------------- C2
def test_format_cell_like_excel():
    assert format_cell(869.065868, "$#,##0") == "$869"
    assert format_cell(0.4895, "0.00%") == "48.95%"
    assert format_cell(2_240_000, '"$"#,##0') == "$2,240,000"
    assert format_cell(1234.5, '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)') == "$1,234.50"
    assert format_cell(89855.4, "#,##0") == "89,855" and format_cell(0.39, "0%") == "39%"
    assert format_cell(datetime(2025, 12, 11), "mmmm d, yyyy") == "December 11, 2025"
    assert format_cell(3.14159, "General") == 3.14159 and format_cell("text", "$#,##0") == "text"


def test_xlsx_to_doc_and_table_tidy(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "LTV", 0.4895
    ws["B1"].number_format = "0.00%"
    ws["A2"], ws["B2"] = "Loan", 2240000.4
    ws["B2"].number_format = '"$"#,##0'
    wb.save(tmp_path / "w.xlsx")
    text = xlsx_to_doc(tmp_path / "w.xlsx", "w.xlsx").text
    assert "R1: A=LTV | B=48.95%" in text and "R2: A=Loan | B=$2,240,000" in text
    t = S.Table(sheet=None, title="Comps", columns=["Address", "Price", "$/SF", "Cap"],
                rows=[["1 Main St", "$1,234.5678", "$869.065868", "6.12345%"]], bold_rows=[])
    assert tidy_table_numbers(t).rows[0] == ["1 Main St", "$1,234.57", "$869", "6.12%"]
    t2 = S.Table(sheet=None, title="Loan Metrics", columns=[], rows=[["Loan PSF:", "$24.929052"], ["LTV:", "39.2982%"]],
                 bold_rows=[])
    assert tidy_table_numbers(t2).rows == [["Loan PSF:", "$25"], ["LTV:", "39.3%"]]
