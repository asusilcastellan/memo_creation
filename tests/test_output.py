"""Output = the .docx + "<memo> - Review.xlsx"; red remarks with IDs that match the workbook; clean cache."""
from __future__ import annotations

import json
import os
import re
import sys
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import load_workbook  # noqa: E402

from memo_app import cache  # noqa: E402
from memo_app.facts import DealFacts  # noqa: E402
from test_offline import docx_text, read_sheet, run  # noqa: E402

REMARK = re.compile(r"\[(REVIEW|MISSING|CHECK) (R\d+): ([^\]]+)\]")


def _highlights(docx_path: Path) -> list[str]:
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    return re.findall(r'<w:highlight w:val="(\w+)"', xml)


def test_two_files_and_clean_run_cache(tmp_path):
    res = run(tmp_path)
    assert sorted(p.name for p in res.memo.parent.iterdir()) == sorted([res.memo.name, res.review.name])
    assert res.review.name == res.memo.stem + " - Review.xlsx"
    run_files = [p for p in cache.runs_dir().rglob("*") if p.is_file()]
    assert [p.name for p in run_files] == ["debug.log"]                    # facts.json / report deleted after success
    assert not list(cache.runs_dir().rglob("*.json")) and not list(cache.runs_dir().rglob("*.md"))


def test_ids_match_between_docx_and_workbook(tmp_path):
    res = run(tmp_path, keep_run_files=True)
    body = docx_text(res.memo)
    assert "Review Notes" not in body
    doc_ids = {m.group(2) for m in REMARK.finditer(body)}
    assert doc_ids and "R1" in doc_ids                                      # e.g. [MISSING R1: insert property photos]
    assert all(len(m.group(3).split()) <= 16 for m in REMARK.finditer(body))
    items = read_sheet(res.review, "Review items", header_row=2)
    wb_ids = {r["ID"] for r in items}
    assert doc_ids <= wb_ids                                                # every remark has a row
    flags = json.loads(next(cache.runs_dir().rglob("facts.json")).read_text(encoding="utf-8"))["facts"]["flags"]
    assert {fl["id"] for fl in flags if fl["inline"]} == doc_ids            # and every inline flag is in the body
    assert {fl["id"] for fl in flags} == wb_ids
    by_id = {r["ID"]: r for r in items}
    for m in REMARK.finditer(body):                                         # same level and title on both sides
        assert by_id[m.group(2)]["Priority"] == m.group(1).title() and by_id[m.group(2)]["Item"] == m.group(3)
    nums = sorted(int(i[1:]) for i in wb_ids)
    assert nums == list(range(1, len(nums) + 1))                            # R1..Rn, inline ones first
    assert max(int(i[1:]) for i in doc_ids) == len(doc_ids)
    assert set(_highlights(res.memo)) == {"red"}


def test_workbook_sheets_status_hyperlinks_and_numbers(tmp_path):
    res = run(tmp_path)
    wb = load_workbook(res.review)
    assert wb.sheetnames == ["Review items", "Online search", "Background checks", "Sources used", "Key figures",
                             "Email chain"]
    ws = wb["Review items"]
    assert "review items - generated" in ws["A1"].value and ws.freeze_panes == "A3" and ws.auto_filter.ref.startswith("A2:I")
    assert [c.value for c in ws[2]] == ["ID", "Priority", "Memo section", "Item", "Details", "Source document",
                                        "Source file(s)", "Status", "Analyst comment"]
    dvs = ws.data_validations.dataValidation
    assert any(dv.type == "list" and "Open,Resolved,Not applicable" in dv.formula1 and "H3" in str(dv.sqref) for dv in dvs)
    assert ws["H3"].value == "Open" and ws["B2"].fill.fgColor.rgb.endswith("002060")
    order = {"Review": 0, "Missing": 1, "Check": 2}
    pri = [order[ws.cell(row=r, column=2).value] for r in range(3, ws.max_row + 1)]
    assert pri == sorted(pri)
    online = wb["Online search"]
    assert online["E2"].hyperlink is not None and online["E2"].hyperlink.target == "https://example.com/smhi-industrial"
    assert online["E2"].value == "SMHI expands in Connecticut"
    figures = {r[0].value: r[1] for r in wb["Key figures"].iter_rows(min_row=2)}
    loan = figures["Castellan loan"]
    assert isinstance(loan.value, (int, float)) and loan.value == 2_240_000 and loan.number_format == '"$"#,##0'
    ltv = figures["LTV (as printed in the memo)"]
    assert isinstance(ltv.value, float) and 0 < ltv.value < 1 and ltv.number_format == "0.0%"
    assert figures["Credit score 1"].value == 793 and figures["Spread"].number_format == "0.00%"
    bg = read_sheet(res.review, "Background checks")
    assert bg[0]["Person"] == "Shaya Margareten" and bg[0]["Role in deal"] == "sponsor" and bg[0]["Credit score"] == 793
    srcs = {r["Document"]: r for r in read_sheet(res.review, "Sources used")}
    assert srcs["Property photos"]["Chosen by"] == "Not found" and srcs["Appraisal"]["Chosen by"] == "Automatic"
    with zipfile.ZipFile(res.review) as z:                                  # nothing technical in the workbook
        xml = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    assert "gpt" not in xml.lower() and "token" not in xml.lower() and str(tmp_path.drive or "C:") + "\\" not in xml
    assert "castellan_memo" not in xml.lower() and "castellanmemo" not in xml.lower()


def test_workbook_failure_keeps_the_memo(tmp_path, monkeypatch):
    import memo_app.pipeline as pipeline

    def boom(*a, **k):
        raise RuntimeError("disk full")
    monkeypatch.setattr(pipeline, "build_review_workbook", boom)
    res = run(tmp_path)
    assert res.memo.exists() and res.review is None
    assert [p.name for p in res.memo.parent.iterdir()] == [res.memo.name]


def test_remark_text_and_linked_flags():
    f = DealFacts(loan_type="acquisition", memo_date="today")
    fl = f.flag("review", "Background check shows adverse items: " + "x " * 40, title="adverse items in background reports",
                section="Sponsorship - Background & credit", source="background", key="background_adverse")
    assert f.remark("background_adverse") == f.marker(fl) and f.find("background_adverse") == [fl]
    fl.id = "R3"
    assert fl.remark_text() == "[REVIEW R3: adverse items in background reports]"
    long = f.flag("check", " ".join(f"w{i}" for i in range(30)))
    assert len(long.title.split()) == 15 and long.details.startswith("w0 w1")


def test_cache_cleanup_and_clear():
    old = time.time() - 10 * 86400
    dl_old, dl_new = cache.downloads_dir() / "a.pdf", cache.downloads_dir() / "b.pdf"
    for p in (dl_old, dl_new):
        p.write_text("x")
    os.utime(dl_old, (old, old))
    run_old = cache.new_run_dir("/Bridge Lending/Current Portfolio/1 Old St")
    (run_old / "debug.log").write_text("x")
    for p in (run_old / "debug.log", run_old):
        os.utime(p, (old, old))
    run_new = cache.new_run_dir("/Bridge Lending/Current Portfolio/2 New St")
    (run_new / "debug.log").write_text("x")
    plan = cache.plan_file("/x/2 New St")
    plan.write_text("{}")

    assert cache.cleanup(7) == 2
    assert not dl_old.exists() and dl_new.exists() and not run_old.exists() and run_new.exists()
    cache.clear()
    assert not any(cache.downloads_dir().iterdir()) and not any(cache.runs_dir().iterdir()) and plan.exists()
    assert cache.deal_slug("/Bridge Lending/Current Portfolio/150 E Aurora St").startswith("150-e-aurora-st-")
