"""Offline end-to-end test: fake deal folder + stub LLM -> real .docx.

Run:  python -m pytest -q tests/test_offline.py      (or: python tests/test_offline.py)
No network, no API keys. Checks file selection, facts/rules, rendering and the Word build.
"""
from __future__ import annotations

import os
import re
import sys
import tempfile
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pymupdf  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from memo_app import schemas as S  # noqa: E402
from memo_app.facts import UserInputs  # noqa: E402
from memo_app.settings import Settings, load_settings  # noqa: E402


# --------------------------------------------------------------------------- fake deal folder
def _pdf(path: Path, text: str, sign: bool = False):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=10)
    if sign:
        p2 = doc.new_page()
        p2.insert_text((72, 72), "DocuSign Envelope ID: 1234\nBORROWER: /s/ Yoichy Nussenzweig", fontsize=10)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


ORIGINAL_ID = "<orig-001@brokerco.com>"
REPLY_ID = "<reply-002@smhi.com>"
DISCLAIMER = ("This e-mail and any attachments are confidential and may be privileged. They are intended solely for the "
              "use of the addressee. If you have received this e-mail in error, please notify the sender immediately and "
              "delete it. Any unauthorized use, disclosure or copying is strictly prohibited.")


def _eml(path: Path, subject: str, sender: str, to: str, when: datetime, body: str, msgid: str,
         in_reply_to: str | None = None, mtime_offset: int = 0) -> Path:
    from email.message import EmailMessage
    from email.utils import format_datetime
    m = EmailMessage()
    m["Subject"], m["From"], m["To"], m["Date"], m["Message-ID"] = subject, sender, to, format_datetime(when), msgid
    if in_reply_to:
        m["In-Reply-To"] = m["References"] = in_reply_to
    m.set_content(body)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(m))
    t = when.timestamp() + mtime_offset
    os.utime(path, (t, t))
    return path


def make_emails(folder: Path) -> None:
    """A broker forward (deal details inside the forward) and the sponsor's reply with quoted history + banner."""
    tz = timezone(timedelta(hours=-4))
    _eml(folder / "Fwd 150 E Aurora St Bridge Loan Request.eml", "Fwd: 150 E Aurora St | Bridge Loan Request",
         "Dan Broker <dan@brokerco.com>", "Analyst <analyst@castellanre.com>", datetime(2025, 10, 1, 9, 0, tzinfo=tz),
         "Please see below - strong sponsor.\n\n---------- Forwarded message ---------\n"
         "From: Shaya Margareten <shaya@smhi.com>\nDate: Tue, Sep 30, 2025\nSubject: 150 E Aurora St | Bridge Loan Request\n\n"
         "We are requesting a $3,290,000 bridge loan (65% LTV) to acquire and renovate 150 E Aurora St, an 89,855 SF "
         "industrial building. We will lease it up and refinance within 24 months.\n\nSent from my iPhone\n",
         ORIGINAL_ID)
    # the reply is written with a NEWER file date, so file order != thread order
    _eml(folder / "RE Fwd 150 E Aurora St.eml", "RE: Fwd: 150 E Aurora St | Bridge Loan Request",
         "Shaya Margareten <shaya@smhi.com>", "Analyst <analyst@castellanre.com>", datetime(2025, 10, 3, 14, 30, tzinfo=tz),
         "CAUTION: This email originated from outside of the organization. Do not click links or open attachments "
         "unless you recognize the sender.\n\nThe roof will be replaced in Q1 2026 and two tenants have signed LOIs.\n\n"
         "Sent from my iPhone\n\n" + DISCLAIMER + "\n\n"
         "On Wed, Oct 1, 2025 at 9:00 AM Dan Broker <dan@brokerco.com> wrote:\n> Please see below - strong sponsor.\n"
         "> We are requesting a $3,290,000 bridge loan\n", REPLY_ID, in_reply_to=ORIGINAL_ID, mtime_offset=10 ** 6)


def make_deal(root: Path) -> Path:
    deal = root / "150 E Aurora St"
    _pdf(deal / "Term Sheet" / "Term Sheet- 150 E Aurora St _v2.pdf", "Blended Loan Amount $2,900,000", sign=True)
    (deal / "Term Sheet").mkdir(parents=True, exist_ok=True)
    import docx
    d = docx.Document(); d.add_paragraph("draft term sheet"); d.save(deal / "Term Sheet" / "Term Sheet- 150 E Aurora St _v2.docx")
    _pdf(deal / "Term Sheet" / "Archive" / "Term Sheet old.pdf", "old")
    _pdf(deal / "PSA" / "DOC-20251104-WA0000.pdf", "whatsapp copy")
    _pdf(deal / "PSA" / "Executed PSA.pdf", "Purchase Price shall be Three Million Dollars ($3,000,000)")
    _pdf(deal / "Resume" / "Shaya bio.pdf", "Shaya Margareten - Founder & Operator, SMHI, Monsey")
    _pdf(deal / "PFS & SREO" / "Personal Financial Statement.pdf", "Net worth 3,680,000")
    _pdf(deal / "PFS & SREO" / "Schedule of real estate.pdf", "REO")
    _pdf(deal / "Background Authorisation & Credit Report" / "Due Diligence Disclosure - Eaglescreen.pdf", "consent")
    _pdf(deal / "Background Authorisation & Credit Report" / "Final Results SMargareten.pdf", "No records")
    _pdf(deal / "Appraisal" / "150 East Aurora Street_Engagement Letter_10.29.25.pdf", "engagement")
    _pdf(deal / "Appraisal" / "2025-3654 - 150 East Aurora Street_Final Appraisal_11.20.25.pdf", "Final appraisal")
    _pdf(deal / "Appraisal" / "2025-3654 - 150 East Aurora Street_Invoice.pdf", "invoice")
    _pdf(deal / "Environmental" / "Phase I ESA - 150 E Aurora.pdf", "No RECs")
    _pdf(deal / "Insurance" / "Evidence of Property Insurance.pdf", "ACORD")
    (deal / "Memo").mkdir(parents=True, exist_ok=True)
    make_emails(deal / "Borrower material X")
    for name in ("150 E Aurora St Model.xlsx", "150 E Aurora St - Net Funding.xlsx"):
        wb = Workbook(); ws = wb.active; ws.title = "Model"; ws["B2"] = "Senior Loan"; ws["C2"] = 2240000; wb.save(deal / name)
    return deal


# --------------------------------------------------------------------------- stub LLM
CANNED = {
    "TermSheet": S.TermSheet(
        property_address="150 E Aurora St, Waterbury, CT 06708", borrower_signed=True, signature_evidence="DocuSign",
        blended_loan_amount=2_900_000, senior_loan_amount=None, junior_lender_name="Reign Funding",
        junior_loan_amount=660_000, loan_purpose="acquisition and renovation", property_description="Approx. 90,000 SF industrial",
        issuing_lender="Castellan Capital", structure="mezz", lead_lender=None, total_loan_amount=None,
        castellan_share_amount=None, castellan_share_pct=None,
        is_floating=True, index_name="1-Month CME SOFR", spread_pct=5.81, floor_rate_pct=9.64, fixed_rate_pct=None,
        rate_cap=None, loan_term_months=24, extension_options=None, origination_fee="2.00%", interest_reserve=None,
        construction_reserve=950_000, construction_reserve_terms="released after borrower spends initial $750,000",
        max_ltv_pct=65, personal_guaranty=True, completion_guaranty=False, carveout_guaranty=False,
        recourse_text="full recourse to borrower and key principals", key_principals=["Shaya Margareten"],
        security_summary="First position mortgage lien; Pledge of cashflow; Personal Guarantees from the Borrower",
        estimated_closing_date="December 11, 2025"),
    "PSA": S.PSA(purchase_price=3_000_000, deposit=50_000, seller="BH EBAY LLC", buyer="TYH 150 AURORA LLC",
                 agreement_date="July 3, 2025", closing_terms=None, property_address="150 E Aurora St", is_executed=True),
    "UWModel": S.UWModel(
        property_address="150 E Aurora St Waterbury, CT 06708", property_type="industrial",
        property_features=["loading docks", "24' ceiling"], year_built="1942", units=1, unit_mix=None,
        gross_sf=90_000, net_rentable_sf=89_855, land_acres=5.0, purchase_price=3_000_000, capex_budget=1_700_000,
        closing_costs=116_000, senior_loan=2_240_000, junior_loan=660_000, sponsor_equity=1_916_000,
        total_sources=4_816_000, total_uses=4_816_000,
        sources=[S.LineItem(label="Senior Loan", amount=2_240_000, pct=47), S.LineItem(label="Junior Loan", amount=660_000, pct=14),
                 S.LineItem(label="Equity", amount=1_916_000, pct=40)],
        uses=[S.LineItem(label="Purchase Price", amount=3_000_000, pct=None), S.LineItem(label="CapEx Budget", amount=1_700_000, pct=None),
              S.LineItem(label="Closing Costs", amount=116_000, pct=None)],
        appraised_value=5_700_000, appraised_value_type="as complete", initial_funding=1_950_000,
        construction_reserve=950_000, facility_leverage=1_568_000, facility_rate_pct=7.30, rate_spread_pct=5.81,
        floor_rate_pct=9.64, index_rate_pct=None, notes=None),
    "PFSResult": S.PFSResult(sponsors=[S.SponsorPFS(full_name="Shaya Margareten", honorific="Mr.", last_name="Margareten",
                                                    as_of_date="2025-09-10", net_worth=3_680_000, cash=150_000,
                                                    marketable_securities=None, real_estate_assets=2_800_000, total_liabilities=None)]),
    "BackgroundResult": S.BackgroundResult(provider="Eaglescreen", people=[S.PersonBackground(
        full_name="Shaya Margareten", credit_score=793, bankruptcies=False, judgments_or_liens=False, criminal_records=False,
        pending_civil_cases=False, past_due_accounts=False, adverse_summary=None)]),
    "ResumeFacts": S.ResumeFacts(people=["Shaya Margareten"], honorifics=["Mr."], company="SMHI", general_contractor=None),
    "Appraisal": S.Appraisal(
        appraiser_firm="Leitner Berman", effective_date="November 2025", property_address="150 E Aurora St, Waterbury, CT 06708",
        property_type="industrial", project_qualifier=None, property_description="Single-story industrial building of 89,855 SF.",
        gross_building_sf=89_855, net_rentable_sf=89_855, units=None, site_acres=5.1, year_built="1942-1976",
        as_is_value=None, as_complete_value=5_700_000, as_stabilized_value=None, gross_sellout_value=None,
        sales_comps=[S.SaleComp(address="1050 Buckingham Street, Watertown, CT", sale_date="Oct-25", sale_price=5_200_000,
                                size_sf=52_241, price_per_sf=100, adjusted_price_per_sf=100),
                     S.SaleComp(address="329 Cooke Street, Plainville, CT", sale_date="Sep-25", sale_price=2_100_000,
                                size_sf=37_671, price_per_sf=56, adjusted_price_per_sf=40)],
        sales_comps_avg_adjusted_psf=70, location="Browns Meadows neighbourhood of Waterbury", access_and_demand_drivers="I-84 and Route 8",
        market_statistics=["Vacancy 4.4% Q4 2025 (CoStar)"], demographics=[], highest_best_use_or_plan=None),
    "Environmental": S.Environmental(provider="CBRE", report_type="Phase I Environmental Site Assessment", phase="phase_i",
                                     report_date="Nov 2025", status="completed", recognized_environmental_conditions=False,
                                     conditions_identified=None, no_further_assessment=None, findings_summary=None,
                                     recommendations=None),
    "EmailInsights": S.EmailInsights(
        loan_purpose="Acquisition and renovation of an industrial building", business_plan="Replace the roof in Q1 2026, lease up",
        exit_strategy="Lease-up then refinance", requested_loan_amount=3_290_000, requested_ltv_or_ltc="65% LTV",
        requested_terms="24 months", property_facts=["89,855 SF industrial building"], sponsor_facts=[],
        timeline=["2025-10-01 - broker forwarded the request"], open_issues=["two tenant LOIs to be provided"],
        key_quotes=["We will lease it up and refinance within 24 months. - Shaya Margareten, 2025-09-30"],
        participants=["Dan Broker (broker, dan@brokerco.com)", "Shaya Margareten (sponsor, shaya@smhi.com)"]),
    "Insurance": S.Insurance(insurer_or_broker="X", document_type="Evidence of property insurance", is_bound=True,
                             property_coverage=4_050_000, liability_per_occurrence=2_000_000, liability_aggregate=None,
                             umbrella=None, builders_risk=None),
}


CANNED["AdverseSearch"] = S.AdverseSearch(
    findings=[S.AdverseFinding(person="Shaya Margareten", finding="Trade press profile of SMHI's industrial acquisitions.",
                               category="press", severity="info", source_url="https://example.com/smhi-industrial",
                               source_title="SMHI expands in Connecticut")],
    overall_assessment="No adverse information found.", adverse_found=False)

CANNED["SectionPlan"] = S.SectionPlan(sentences=[
    S.PlannedSentence(purpose="state the loan amount and purpose", facts_used=["$2,240,000", "acquisition"]),
    S.PlannedSentence(purpose="describe the property", facts_used=["150 E Aurora St, Waterbury, CT 06708", "89,855 SF"]),
])

# Plausible section texts (lengths inside the playbook targets, numbers only from the facts / sources).
STUB_TEXT = {
    "transaction_overview": (
        " for the acquisition and renovation of the industrial property located at 150 E Aurora St, Waterbury, CT."
        " The subject property is a single-story industrial building totaling 89,855 square feet. Reign Funding will"
        " provide a $660,000 junior loan subordinate to Castellan’s position. The loan will fund part of the $3,000,000"
        " purchase price and of the $1,700,000 capital expenditure budget, and the borrower plans to lease up the"
        " property and repay Castellan through a refinance with a permanent loan."),
    "business_background": (
        "According to the resume provided to Castellan, Shaya Margareten is the founder & operator of SMHI, a real"
        " estate investment company based in Monsey, New York. Mr. Margareten acquires, renovates and manages"
        " industrial and residential properties across the region and oversees acquisitions, construction management"
        " and leasing for the company. His experience includes repositioning older industrial buildings and leasing"
        " them to local and regional tenants, which is directly relevant to the business plan for the subject property."),
    "submarket_overview": (
        "The subject property is located in the Browns Meadows neighbourhood of Waterbury, Connecticut, an established"
        " industrial and residential area close to the city center. The property benefits from direct access to I-84"
        " and Route 8, which connect Waterbury with Hartford, New Haven and the wider New England distribution network."
        " Based on the appraisal report, the area is served by a deep base of manufacturing, logistics and healthcare"
        " employers that support steady demand for industrial space. Based on the data provided by CoStar, vacancy in"
        " the submarket was 4.4% as of the fourth quarter of 2025, which supports the borrower’s plan to lease up the"
        " property after the renovation."),
    "business_plan": (
        "The borrower is acquiring the property for $3,000,000 and has budgeted $1,700,000 for capital expenditures"
        " to modernize the building, including roof and loading improvements. The renovation will be funded in part"
        " through a construction reserve released once the borrower has spent its initial equity. Based on the"
        " appraisal report, the as-complete value of the property is $5,700,000. Upon completion, the borrower plans"
        " to lease the building to industrial tenants and refinance Castellan’s loan with a permanent loan."),
    "sponsor_experience": (
        "Based on the resume provided by Mr. Margareten, he has substantial experience as an owner and operator of"
        " industrial and residential real estate in the region."),
}


class StubLLM:
    last: "StubLLM | None" = None

    def __init__(self, settings, log=print):
        self.usage = {"calls": 0, "input_tokens": 0, "output_tokens": 0}
        self.prompts = []
        self.extract_inputs: dict[str, str] = {}
        StubLLM.last = self

    def extract(self, schema, instructions, text=None, images_b64=None):
        self.usage["calls"] += 1
        self.extract_inputs[schema.__name__] = text or ""
        return CANNED[schema.__name__]

    def write(self, system, user):
        self.usage["calls"] += 1
        self.prompts.append(user)
        marker = "begin the paragraph with exactly these words and continue the sentence:\n"
        if marker in user:
            opening = user.split(marker, 1)[1].splitlines()[0]
            return opening + STUB_TEXT["transaction_overview"]
        section = re.search(r"Write the (.+?) for a", user).group(1).lower().replace(" ", "_")
        return STUB_TEXT[section]

    def web_search(self, schema, prompt):
        return CANNED["AdverseSearch"]


@contextmanager
def offline_rates():
    """No network in tests: the NY Fed / FRED fetchers return fixed values (restored afterwards)."""
    from memo_app import rates
    stubs = {"fetch_sofr": lambda *a, **k: 4.30, "fetch_sofr_30d_avg": lambda *a, **k: 4.25,
             "fetch_prime_fred": lambda *a, **k: None}
    saved = {k: getattr(rates, k) for k in stubs}
    for k, v in stubs.items():
        setattr(rates, k, v)
    try:
        yield
    finally:
        for k, v in saved.items():
            setattr(rates, k, v)


def run(tmp: Path, **debug):
    """-> pipeline.MemoOutput (.memo = the .docx, .review = the review workbook)."""
    with offline_rates():
        return _run(tmp, **debug)


def _run(tmp: Path, **debug):
    import memo_app.pipeline as pipeline
    pipeline.LLM = StubLLM
    settings: Settings = load_settings()
    settings.openai_api_key = "stub"
    settings.raw.setdefault("debug", {}).update(debug)
    deal = make_deal(tmp)
    ui = UserInputs(deal_path=str(deal), loan_type="acquisition", notes="Borrower will lease up and refinance.",
                    run_web_search=True, output_dir=str(tmp / "out"))
    return pipeline.generate(ui, settings)


def docx_text(path: Path) -> str:
    import docx
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)


def read_sheet(path: Path, name: str, header_row: int = 1) -> list[dict]:
    from openpyxl import load_workbook
    ws = load_workbook(path)[name]
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def test_end_to_end(tmp_path):
    res = run(tmp_path, keep_run_files=True)
    out = res.memo
    assert out.exists() and res.review and res.review.exists()
    assert sorted(p.name for p in out.parent.iterdir()) == sorted([out.name, res.review.name])
    body = docx_text(out)
    assert "Review Notes" not in body
    assert "Castellan Capital will provide a senior loan of $2,240,000" in body
    assert "39%" in body and "Leitner Berman" in body and "Reign Funding" in body
    assert "Mr. Margareten will provide a full personal guaranty" in body
    # term sheet index 1-Month CME SOFR, no value typed -> overnight SOFR proxy 4.30% + 5.81% = 10.11% (> floor)
    assert "floating based on the 1-Month CME Term SOFR index plus a spread of 5.81%" in body
    assert "effective rate at closing is expected to be 10.11%" in body
    # the review workbook carries the details
    items = read_sheet(res.review, "Review items", header_row=2)
    details = " | ".join(f"{r['Item']} :: {r['Details']}" for r in items)
    assert "used overnight SOFR 4.30% from the NY Fed as a proxy" in details
    assert "Email chain requested loan $3,290,000 differs from the term sheet loan $2,240,000" in details
    assert "Email chain requested 65% LTV; the computed LTV is 39%." in details
    sources = {r["Document"]: r for r in read_sheet(res.review, "Sources used")}
    assert sources["Signed term sheet"]["File(s)"] == "Term Sheet/Term Sheet- 150 E Aurora St _v2.pdf"
    assert "Executed PSA.pdf" in sources["Purchase agreement (PSA)"]["File(s)"]
    assert "Final Appraisal" in sources["Appraisal"]["File(s)"]
    figures = {r["Figure"]: r for r in read_sheet(res.review, "Key figures")}
    assert abs(figures["Index rate"]["Value"] - 0.043) < 1e-9 and "proxy" in figures["Index rate"]["Source document"]
    # email chain: both .eml files read in date order, conflicts listed with them
    chain = read_sheet(res.review, "Email chain")
    subjects = [r["Subject"] for r in chain]
    assert subjects[0].startswith("Fwd: 150 E Aurora St") and subjects[1].startswith("RE: Fwd: 150 E Aurora St")
    assert sum(1 for r in chain if (r["Subject"] or "").startswith("Conflict with the term sheet")) == 2
    stub = StubLLM.last
    sent = stub.extract_inputs["EmailInsights"]
    assert "We are requesting a $3,290,000 bridge loan" in sent        # forwarded content kept
    assert "CAUTION" not in sent and "> Please see below" not in sent and "Sent from my iPhone" not in sent
    # the email facts reach the writer as background facts (transaction overview + business plan), never as a
    # label the model could echo; every section passed the style check without a repair call
    assert sum("BACKGROUND FACTS ON PURPOSE, PLAN AND EXIT" in p for p in stub.prompts) == 2
    assert not any("EMAIL CHAIN" in p.upper() for p in stub.prompts)
    assert any("EXIT STRATEGY (background fact): Lease-up then refinance" in p for p in stub.prompts)
    assert len(stub.prompts) == 5 and "style check" not in details and "email chain" not in body.lower()
    import json
    from memo_app import cache
    fj = next(cache.runs_dir().rglob("facts.json"))                     # kept because keep_run_files=True
    msgs = json.loads(fj.read_text(encoding="utf-8"))["email_chain"]["messages"]
    assert [m["message_id"] for m in msgs] == [ORIGINAL_ID, REPLY_ID] and "body_text" not in msgs[0]


if __name__ == "__main__":
    keep = Path(os.environ.get("KEEP_DIR", tempfile.mkdtemp()))
    print("output:", run(keep).memo)
